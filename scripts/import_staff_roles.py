"""Preview staff-register roles and memberships; write only with --apply."""

from __future__ import annotations

import argparse
import asyncio
import json
import logging
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from urllib.parse import quote

from openpyxl import load_workbook
from pydantic import StrictStr, TypeAdapter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from config import settings
from doc_standards.roles import RoleBook, norm, similarity
from graph import client as graph_client
from graph.exceptions import GraphNotFoundError
from groups import service as groups_service
from scripts.import_roles_from_jds import _table, read_existing_roles


@dataclass(frozen=True)
class StaffPerson:
    name: str
    email: str = field(repr=False)
    title: str
    department: str = ""


@dataclass
class RoleProposal:
    name: str
    departments: set[str] = field(default_factory=set)
    reviewed: bool = False
    near: list[tuple[str, int]] = field(default_factory=list)

    @property
    def create(self) -> bool:
        return self.reviewed or not self.near


@dataclass
class StaffPlan:
    staff: list[StaffPerson]
    existing: list[dict]
    proposals: list[RoleProposal]
    people: list[tuple[str, str, str]] = field(default_factory=list)
    additions: dict[str, list[dict]] = field(default_factory=lambda: defaultdict(list), repr=False)
    already: Counter = field(default_factory=Counter)
    unmatched: list[tuple[str, str, str, str]] = field(default_factory=list)
    unresolved: list[tuple[str, str, str]] = field(default_factory=list)
    missing_mappings: list[tuple[str, str, str]] = field(default_factory=list)


def split_title(title: str) -> list[str]:
    """Only slash separates roles; 'and' and '&' belong to a role's name."""
    return list(dict.fromkeys(part.strip() for part in title.split("/") if part.strip()))


def read_staff_register(path: Path) -> list[StaffPerson]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Staff List"]
        rows = sheet.iter_rows(min_row=2, values_only=True)
        headers = {str(value).strip(): i for i, value in enumerate(next(rows)) if value}
        required = {"S/N", "Staff Name", "Email", "Job Title", "Department"}
        if not required <= headers.keys():
            raise ValueError("Staff List row 2 is missing required columns.")
        staff = []
        for row in rows:
            def cell(key):
                value = row[headers[key]]
                return str(value) if value is not None else ""
            if not cell("S/N").strip() or not cell("Staff Name").strip():
                continue
            # Keep the job title verbatim: mapping keys are the register's exact text.
            staff.append(StaffPerson(cell("Staff Name").strip(), cell("Email").strip(),
                                     cell("Job Title"), cell("Department").strip()))
        return staff
    finally:
        workbook.close()


def load_mapping(path: Path) -> dict[str, list[str]]:
    raw = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(raw, dict):
        raise ValueError("The mapping must be a JSON object.")
    mapping = TypeAdapter(dict[str, list[StrictStr]]).validate_python(
        {key: value for key, value in raw.items() if not key.startswith("_")}
    )
    for title, targets in mapping.items():
        if not title.strip() or not targets or any(not target.strip() for target in targets):
            raise ValueError("Mapping titles and target role names must not be empty.")
        mapping[title] = list(dict.fromkeys(target.strip() for target in targets))
    return mapping


def near_matches(name: str, roles: list[dict]) -> list[tuple[str, int]]:
    matches = []
    for role in roles:
        if norm(name) == norm(role["name"]):
            continue
        score = max(similarity(name, spelling)[0]
                    for spelling in [role["name"], *role.get("aliases", [])])
        if score >= 80:
            matches.append((role["name"], score))
    return sorted(set(matches), key=lambda match: (-match[1], match[0]))


def propose_roles(staff: list[StaffPerson], mapping: dict[str, list[str]], existing: list[dict]) -> list[RoleProposal]:
    book = RoleBook.from_groups(existing)
    proposals: dict[str, RoleProposal] = {}
    for person in staff:
        targets = mapping.get(person.title)
        # A title is ONE role unless the reviewed mapping says otherwise: a slash can belong to the
        # role's own name ("Finance/Admin Manager"), so it is never split on a guess.
        for part in (targets if targets is not None else [person.title]):
            if book.find(part):
                continue
            # A mapped role is created only if it is the title itself or one of its slash parts: a typo in
            # the mapping file is reported (missing target), never turned into a new role.
            reviewed = targets is not None
            if reviewed and norm(part) not in {norm(person.title), *map(norm, split_title(person.title))}:
                continue
            proposal = proposals.setdefault(norm(part), RoleProposal(part))
            if person.department:
                proposal.departments.add(person.department)
            proposal.reviewed |= reviewed
    for key, proposal in proposals.items():
        other_proposals = [{"name": other.name} for other_key, other in proposals.items() if other_key != key]
        proposal.near = near_matches(proposal.name, [*existing, *other_proposals])
    return sorted(proposals.values(), key=lambda proposal: proposal.name.casefold())


async def resolve_person(email: str) -> tuple[dict | None, str]:
    """The /grc/users/resolve lookup pattern, with exact, unique email matching."""
    if not email or "@" not in email:
        return None, "Missing or invalid email in register"
    try:
        data = await graph_client._request(
            "GET", f"{settings.graph_base_url}/users/{quote(email, safe='')}",
            params={"$select": "id,displayName,mail,userPrincipalName"}, context="Resolve staff account",
        )
        results = [data]
    except GraphNotFoundError:
        escaped = email.replace("'", "''")
        data = await graph_client._request(
            "GET", f"{settings.graph_base_url}/users",
            params={"$filter": f"mail eq '{escaped}' or userPrincipalName eq '{escaped}'",
                    "$select": "id,displayName,mail,userPrincipalName"}, context="Find staff account",
        )
        results = data.get("value", [])
    exact = {user["id"]: user for user in results if user.get("id") and
             email.casefold() in {(user.get("mail") or "").casefold(),
                                  (user.get("userPrincipalName") or "").casefold()}}
    if len(exact) != 1:
        return None, "Multiple matching Entra accounts" if exact else "No matching Entra account"
    oid, data = next(iter(exact.items()))
    profile = await graph_client.resolve_user(oid)
    return {"oid": oid, "display_name": profile.get("display_name") or data.get("displayName") or "",
            "email": data.get("mail") or data.get("userPrincipalName") or email}, ""


async def build_plan(staff: list[StaffPerson], mapping: dict[str, list[str]]) -> StaffPlan:
    existing = await read_existing_roles()
    proposals = propose_roles(staff, mapping, existing)
    planned = [{"id": f"new:{norm(p.name)}", "name": p.name, "members": []}
               for p in proposals if p.create]
    future = [*existing, *planned]
    book = RoleBook.from_groups(future)
    current = RoleBook.from_groups(existing)
    by_id = {str(role["id"]): role for role in future}
    plan = StaffPlan(staff, existing, proposals)

    for title, targets in mapping.items():
        for target in targets:
            found = current.find(target)
            if not found or by_id[found.group_id].get("status") == "Withdrawn":
                pending = book.find(target)
                state = "Will create" if pending and pending.group_id.startswith("new:") else "Missing or withdrawn; skipped"
                plan.missing_mappings.append((title, target, state))

    semaphore = asyncio.Semaphore(6)
    async def lookup(email):
        async with semaphore:
            return await resolve_person(email)
    emails = list(dict.fromkeys(person.email.casefold() for person in staff))
    accounts = dict(zip(emails, await asyncio.gather(*(lookup(email) for email in emails))))
    counts = Counter(person.title for person in staff)
    seen_titles, seen_memberships = set(), set()
    for person in staff:
        user, reason = accounts[person.email.casefold()]
        if not user:
            plan.unresolved.append((person.name, person.title, reason))
        targets = mapping.get(person.title)
        exact = book.find(person.title)
        if targets is None and exact:
            targets = [exact.name]
        if targets is None:
            if person.title not in seen_titles:
                suggestions = current.suggest(person.title, limit=1)
                closest = suggestions[0] if suggestions else None
                plan.unmatched.append((person.title or "(blank title)", str(counts[person.title]),
                                       closest.role.name if closest else "—", str(closest.score) if closest else "—"))
                seen_titles.add(person.title)
            plan.people.append((person.name, person.title or "—", "Unmatched title; no membership"))
            continue
        for target in targets:
            role = book.find(target)
            if not role or by_id[role.group_id].get("status") == "Withdrawn":
                plan.people.append((person.name, target, "Missing or withdrawn role; skipped"))
                continue
            if not user:
                plan.people.append((person.name, role.name, reason))
                continue
            key = (role.group_id, user["oid"])
            if key in seen_memberships:
                continue
            seen_memberships.add(key)
            present = any(member.get("oid") == user["oid"] for member in by_id[role.group_id].get("members", []))
            if present:
                plan.already[role.name] += 1
            else:
                plan.additions[role.name].append({**user, "display_name": user["display_name"] or person.name})
            plan.people.append((person.name, role.name, "Already a member" if present else "Add after role creation" if role.group_id.startswith("new:") else "Add"))
    return plan


def print_plan(plan: StaffPlan, *, apply: bool) -> None:
    _table(("Mode", "Staff", "Existing roles", "New roles", "Roles needing review", "Membership additions", "Already members"),
           [("APPLY" if apply else "DRY RUN", str(len(plan.staff)), str(len(plan.existing)),
             str(sum(p.create for p in plan.proposals)), str(sum(not p.create for p in plan.proposals)),
             str(sum(len(users) for users in plan.additions.values())), str(sum(plan.already.values())))])
    print("\nProposed roles (near-duplicates stay pending unless explicitly confirmed by the mapping):")
    _table(("Role", "Category", "Action"), [(p.name, "; ".join(sorted(p.departments)) or "—",
            "Create (reviewed mapping)" if p.reviewed else "Create" if p.create else "Review; no creation") for p in plan.proposals])
    print("\nNear-duplicates (existing roles and other register titles; suggestions only):")
    _table(("Proposed role", "Similar role", "Score"),
           [(p.name, name, str(score)) for p in plan.proposals for name, score in p.near])
    print("\nMembership counts per role:")
    _table(("Role", "Add", "Already members"),
           [(name, str(len(plan.additions.get(name, []))), str(plan.already[name]))
            for name in sorted(set(plan.additions) | set(plan.already))])
    print("\nPeople and roles:")
    _table(("Person", "Role", "Action"), plan.people)
    print(f"\nUnmatched titles ({len(plan.unmatched)} distinct; no memberships):")
    _table(("Title", "People", "Closest existing role", "Score"), plan.unmatched)
    print(f"\nPeople without a unique Entra account ({len(plan.unresolved)}; skipped):")
    _table(("Person", "Title", "Reason"), plan.unresolved)
    print(f"\nMapping targets not currently active ({len(plan.missing_mappings)}):")
    _table(("Register title", "Target role", "Action"), plan.missing_mappings)


async def apply_plan(plan: StaffPlan) -> tuple[int, int]:
    created, added = 0, 0
    known = list(plan.existing)
    for proposal in plan.proposals:
        if not proposal.create:
            continue
        latest = await read_existing_roles()
        if RoleBook.from_groups([*known, *latest]).find(proposal.name):
            continue
        # A newly appeared near-duplicate needs review too, before any write.
        if not proposal.reviewed and near_matches(proposal.name, latest):
            _table(("Role", "Action"), [(proposal.name, "New near-duplicate; skipped for review")])
            continue
        role = await groups_service.create_group(
            name=proposal.name, description="Imported from staff-register job titles.",
            category="; ".join(sorted(proposal.departments)), members=[],
            creator_oid="", creator_name="Staff role import", aliases=[],
        )
        known.append(role)
        created += 1
    for name, people in plan.additions.items():
        latest = await read_existing_roles()
        found = RoleBook.from_groups(latest).find(name)
        role = next((role for role in latest if found and str(role["id"]) == found.group_id), None)
        if not role or role.get("status") == "Withdrawn":
            _table(("Role", "Action"), [(name, "Role unavailable; memberships skipped")])
            continue
        members = role.get("members", [])
        if any(not member.get("oid") for member in members):
            _table(("Role", "Action"), [(name, "Existing member entry needs repair; skipped")])
            continue
        existing_oids = {member["oid"] for member in members}
        new_people = [person for person in people if person["oid"] not in existing_oids]
        if new_people:
            await groups_service.add_members(role["id"], new_people)
            added += len(new_people)
    _table(("Roles created", "Memberships added"), [(str(created), str(added))])
    return created, added


async def import_staff(staff: list[StaffPerson], mapping: dict[str, list[str]], *, apply: bool = False) -> StaffPlan:
    plan = await build_plan(staff, mapping)
    print_plan(plan, apply=apply)
    if apply:
        await apply_plan(plan)
    return plan


async def _run(args: argparse.Namespace) -> None:
    staff, mapping = await asyncio.gather(asyncio.to_thread(read_staff_register, args.register_path),
                                          asyncio.to_thread(load_mapping, args.mapping))
    if not staff:
        raise ValueError("The staff register contains no staff rows.")
    await graph_client.startup()
    try:
        await import_staff(staff, mapping, apply=args.apply)
    finally:
        await graph_client.shutdown()


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("register_path", type=Path, help="Path to the staff-register XLSX")
    parser.add_argument("--mapping", type=Path, default=Path(__file__).with_name("staff_role_map.json"))
    parser.add_argument("--apply", action="store_true", help="Create approved missing roles and add eligible members")
    args = parser.parse_args(argv)
    # Graph errors and httpx URLs can contain email addresses. Only the tables are printed.
    previous_logging = logging.root.manager.disable
    logging.disable(logging.CRITICAL)
    try:
        asyncio.run(_run(args))
    except Exception as exc:
        _table(("Import error",), [(f"{type(exc).__name__}: import stopped; no further writes attempted.",)])
        return 1
    finally:
        logging.disable(previous_logging)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
