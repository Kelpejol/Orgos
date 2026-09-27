#!/usr/bin/env python3
# =============================================================================
# scripts/build_cdt_master_templates.py — generate the CDT v06 master templates
#
# Builds one master .docx per document type (Policy, Procedure, Combined,
# Manual, Guideline, Standard, SLA — the exact set the frontend's "New
# document" form and the AI drafter offer) and, optionally, uploads them to
# the ORGOS LIBRARY's `Templates/` folder where lifecycle.templates expects
# them (`Templates/{doc_type}.docx`).
#
# Shape is the one confirmed by the compliance officer (docs/CDT-V06-TEMPLATE-
# ANALYSIS.md §8): a 13-row cover control table with a growing `standards`
# tail, then a revision-history table — both using docxtpl's 3-row loop
# pattern (dedicated for-row / data-row / endfor-row; combining the for-tag
# with data in one row silently produces ZERO repeated rows — see
# tests/lifecycle/test_merge.py). Blank header; fixed footer (no per-document
# markers there — confirmed retired).
#
# Styling (logo, fonts, sizes, colours, table shading, footer field structure)
# was reverse-engineered directly from a real, already-issued Dragnet document
# ("docs/Acceptable Use Standard- V2.docx") by unzipping its OOXML and reading
# the actual run/cell properties — not guessed. The extracted logo image
# lives at assets/branding/dragnet_logo.png. See docs/CDT-V06-TEMPLATE-
# ANALYSIS.md for the full extracted spec.
#
# Every generated template is validated with lifecycle.templates.
# validate_template() before it is written/uploaded — a template that fails
# structural validation is never produced silently.
#
# Usage:
#   python scripts/build_cdt_master_templates.py                # dry run — validates, saves locally to ./cdt_templates_out/
#   python scripts/build_cdt_master_templates.py --upload        # also uploads to ORGOS LIBRARY/Templates/
#   python scripts/build_cdt_master_templates.py --upload --only Procedure,Policy
# =============================================================================

import argparse
import asyncio
import io
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Emu, Inches, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from lifecycle.templates import template_filename_for, validate_template

# The exact 7 document types offered by the frontend's "New document" form
# and the AI drafter (frontend/src/pages/DocumentLifecycle/index.jsx
# DOC_TYPES; agents/policy_drafter/service.py TYPE_CODES).
DOCUMENT_TYPES = ["Policy", "Procedure", "Combined", "Manual", "Guideline", "Standard", "SLA"]

# ── Exact styling extracted from docs/Acceptable Use Standard- V2.docx ──────
FONT_NAME = "Cambria"

TITLE_COLOR      = RGBColor(0x1A, 0x1A, 0x1A)   # cover title
BLACK            = RGBColor(0x00, 0x00, 0x00)   # body / revision-heading / values
WHITE            = RGBColor(0xFF, 0xFF, 0xFF)   # history table header text
COVER_LABEL_FILL = "EEF1F6"                      # cover table label-cell shading
HISTORY_HEAD_FILL = "1F3864"                     # revision-history header row (dark navy)

LOGO_PATH  = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                          "assets", "branding", "dragnet_logo.png")
LOGO_WIDTH = Inches(1.54)

# A4 page + the exact (slightly non-round) margins found in the source file.
PAGE_WIDTH  = Emu(7560310)
PAGE_HEIGHT = Emu(10692130)
PAGE_MARGIN = Emu(635000)

# Cover table column widths (label / value), in EMU, from the source file.
COVER_COL_WIDTHS = (Emu(1597025), Emu(4134485))
# Revision-history table column widths (Version / Date / Purpose / Approved By).
HISTORY_COL_WIDTHS = (Emu(606425), Emu(690245), Emu(3576320), Emu(858520))

# The 13 fixed cover rows, in the confirmed order. (label, marker)
COVER_ROWS = [
    ("Document Title",        "{{ title }}"),
    ("Document Code",         "{{ doc_code }}"),
    ("Document Type / Layer", "{{ type_layer }}"),
    ("Domain",                "{{ domain }}"),
    ("Parent Document",       "{{ parent_document }}"),
    ("Classification",        "{{ classification }}"),
    ("Distribution",          "{{ distribution }}"),
    ("Owner",                 "{{ owner }}"),
    ("First Pass Approval",   "{{ first_pass_approval }}"),
    ("Final Approval",        "{{ final_approval }}"),
    ("Version",               "{{ version }}"),
    ("Effective Date",        "{{ effective_date }}"),
    ("Next Review Due",       "{{ next_review_due }}"),
]

HISTORY_HEADERS = ["Version", "Date", "Purpose of Change", "Approved By"]


def _run(para, text, *, bold=False, italic=False, size_pt=11, color=None, font_name=FONT_NAME):
    # The source document's Normal style has NO spacing/line-height override
    # (docDefaults leaves pPrDefault empty) — zero it here too, or every
    # table row ends up visibly taller than the source's tightly-packed rows.
    pf = para.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    run = para.add_run(text)
    run.bold, run.italic = bold, italic
    run.font.size = Pt(size_pt)
    run.font.name = font_name
    if color:
        run.font.color.rgb = color
    return run


def _set_table_borders(table) -> None:
    """
    The source document does NOT reference the built-in "Table Grid" style —
    it sets w:tblBorders directly (single, sz=4, auto colour) with no named
    style at all. Using "Table Grid" instead pulls in the renderer's own
    built-in cell-margin defaults, which are wider than the source's and
    caused header text (e.g. "Version") to wrap that shouldn't. Match the
    source exactly: explicit borders, no named table style.
    """
    tblPr = table._tbl.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        table._tbl.insert(0, tblPr)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        borders.append(el)
    tblPr.append(borders)

    # Match the source's effective cell padding (inherited there from the
    # built-in "Normal Table" style: 0 top/bottom, 108 dxa left/right).
    # Leaving this unset lets the renderer fall back to ITS OWN default
    # margins, which are wider — that was making rows taller and even
    # wrapping short header words like "Version".
    cell_mar = OxmlElement("w:tblCellMar")
    for edge, width in (("top", "0"), ("left", "108"), ("bottom", "0"), ("right", "108")):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), width)
        el.set(qn("w:type"), "dxa")
        cell_mar.append(el)
    tblPr.append(cell_mar)


def _shade_cell(cell, fill_hex: str) -> None:
    """Set a table cell's background shading (w:shd) to the given hex fill."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tcPr.append(shd)


def _set_col_widths(table, widths) -> None:
    """python-docx needs the width set on every cell in a column, not just
    the Column object, or Word re-flows to equal widths on open."""
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width
    for col, width in zip(table.columns, widths):
        col.width = width


def _add_field(paragraph, instr: str, *, bold: bool, size_pt: int, cached_text: str = ""):
    """
    Build a real Word field (begin/instrText/separate/cached-result/end) —
    e.g. PAGE or NUMPAGES — matching the exact run-property shape found in
    the source document's footer, so Word displays a live value instead of
    the literal field code.
    """
    def _field_run(tag=None, text=None, instr_text=None):
        run = paragraph.add_run()
        run.bold = bold
        run.font.size = Pt(size_pt)
        run.font.name = FONT_NAME
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
            run._r.append(el)
        elif instr_text is not None:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = instr_text
            run._r.append(el)
        elif text is not None:
            run.add_text(text)
        return run

    _field_run(tag="begin")
    _field_run(instr_text=f" {instr} ")
    _field_run(tag="separate")
    _field_run(text=cached_text or "1")
    _field_run(tag="end")


def _add_loop_rows(table, *, loop_var: str, iterable: str, cols: list[str], label_fill: str = ""):
    """
    The docxtpl row-loop pattern that ACTUALLY repeats a row: a dedicated
    for-row (tag alone), a dedicated data-row (the {{ }} markers), and a
    dedicated endfor-row (tag alone). Putting the for-tag in the same row as
    data markers produces zero output rows — verified while building the CDT
    test suite; do not "simplify" this to two rows.
    """
    for_row = table.add_row()
    _run(for_row.cells[0].paragraphs[0], f"{{%tr for {loop_var} in {iterable} %}}", size_pt=10)
    data_row = table.add_row()
    for i, field in enumerate(cols):
        if i == 0 and label_fill:
            _shade_cell(data_row.cells[i], label_fill)
        _run(data_row.cells[i].paragraphs[0], f"{{{{ {loop_var}.{field} }}}}", size_pt=10,
             bold=(i == 0 and bool(label_fill)), color=BLACK)
    end_row = table.add_row()
    _run(end_row.cells[0].paragraphs[0], "{%tr endfor %}", size_pt=10)


def build_master_template(doc_type: str) -> bytes:
    """Build one master template's bytes for the given document type label."""
    doc = Document()

    # ── Page geometry — matches the source document (A4 + its exact margins) ──
    section = doc.sections[0]
    section.page_width, section.page_height = PAGE_WIDTH, PAGE_HEIGHT
    section.left_margin = section.right_margin = PAGE_MARGIN
    section.top_margin = section.bottom_margin = PAGE_MARGIN

    # Document-wide default font (Cambria 11pt), so any un-styled run (e.g.
    # the AI body content appended after this template) matches too.
    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(11)

    # ── Brand block: real Dragnet logo, then the wordmark text beneath it ───
    logo_p = doc.add_paragraph()
    logo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo_p.add_run().add_picture(LOGO_PATH, width=LOGO_WIDTH)

    brand = doc.add_paragraph()
    brand.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(brand, "DRAGNET SOLUTIONS", bold=True, size_pt=13)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(subtitle, "{{ domain }} · {{ doc_type }}", size_pt=9)
    doc.add_paragraph()

    # ── Big bold document title ──────────────────────────────────────────────
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = _run(title_p, "{{ title }}", bold=True, size_pt=20, color=TITLE_COLOR)
    title_run.font.all_caps = True  # source document displays the title in full caps
    doc.add_paragraph()

    # ── Cover control table (13 fixed rows + growing standards tail) ────────
    cover = doc.add_table(rows=0, cols=2)
    _set_table_borders(cover)
    for label, marker in COVER_ROWS:
        row = cover.add_row()
        _shade_cell(row.cells[0], COVER_LABEL_FILL)
        row.cells[0].text = ""
        _run(row.cells[0].paragraphs[0], label, bold=True, size_pt=10)
        row.cells[1].text = ""
        _run(row.cells[1].paragraphs[0], marker, size_pt=10, color=BLACK)
    _add_loop_rows(cover, loop_var="s", iterable="standards", cols=["name", "reference"],
                   label_fill=COVER_LABEL_FILL)
    _set_col_widths(cover, COVER_COL_WIDTHS)

    doc.add_paragraph()

    # ── Revision history ──────────────────────────────────────────────────────
    hist_heading = doc.add_paragraph()
    _run(hist_heading, "REVISION HISTORY", bold=True, size_pt=12, color=BLACK)
    hist = doc.add_table(rows=1, cols=4)
    _set_table_borders(hist)
    for i, h in enumerate(HISTORY_HEADERS):
        _shade_cell(hist.rows[0].cells[i], HISTORY_HEAD_FILL)
        hist.rows[0].cells[i].text = ""
        _run(hist.rows[0].cells[i].paragraphs[0], h, bold=True, size_pt=10, color=WHITE)
    _add_loop_rows(hist, loop_var="r", iterable="revision_history",
                   cols=["version", "date", "purpose", "approved_by"])
    _set_col_widths(hist, HISTORY_COL_WIDTHS)

    # ── Footer: "Page {PAGE} of {NUMPAGES}", right-aligned, real Word fields ──
    footer = section.footer
    fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _run(fp, "Page ", size_pt=8)
    _add_field(fp, "PAGE", bold=True, size_pt=9)
    _run(fp, " of ", size_pt=8)
    _add_field(fp, "NUMPAGES", bold=True, size_pt=9)

    # ── Page break so the body always starts fresh, however many revision
    # rows or standards rows the register ends up growing to. ────────────────
    doc.add_page_break()
    body_start = doc.add_paragraph()
    _run(body_start, "[Body content begins here — 1.0 Purpose & Scope, etc.]",
         italic=True, size_pt=9)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def _upload(doc_type: str, content: bytes) -> str:
    from graph.client import ensure_drive_folder, resolve_compliance_drive, upload_bytes_to_drive
    from config import settings

    _, drive_id = await resolve_compliance_drive()
    await ensure_drive_folder(drive_id, settings.cdt_templates_folder)
    filename = template_filename_for(doc_type)
    result = await upload_bytes_to_drive(drive_id, settings.cdt_templates_folder, filename, content)
    return result.get("webUrl", "")


async def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--upload", action="store_true",
                        help="Upload to the ORGOS LIBRARY Templates/ folder (default: local dry-run only)")
    parser.add_argument("--only", default="",
                        help="Comma-separated subset of document types to build, e.g. 'Procedure,Policy'")
    parser.add_argument("--out-dir", default="cdt_templates_out",
                        help="Local directory to save generated templates (dry-run and as a local copy either way)")
    args = parser.parse_args()

    types = [t.strip() for t in args.only.split(",") if t.strip()] if args.only else DOCUMENT_TYPES
    unknown = [t for t in types if t not in DOCUMENT_TYPES]
    if unknown:
        print(f"Unknown document type(s): {unknown}. Valid: {DOCUMENT_TYPES}")
        return 1

    os.makedirs(args.out_dir, exist_ok=True)

    if args.upload:
        from graph import client as gc
        await gc.startup()

    failed = []
    try:
        for doc_type in types:
            print(f"\n=== {doc_type} ===")
            content = build_master_template(doc_type)

            result = validate_template(content)
            print(f"  validation: {result.summary()}")
            if not result.ok:
                failed.append(doc_type)
                continue

            local_path = os.path.join(args.out_dir, template_filename_for(doc_type))
            with open(local_path, "wb") as f:
                f.write(content)
            print(f"  saved locally: {local_path}")

            if args.upload:
                url = await _upload(doc_type, content)
                print(f"  uploaded: {url}")
    finally:
        if args.upload:
            from graph import client as gc
            await gc.shutdown()

    print(f"\n{'='*60}")
    if failed:
        print(f"FAILED validation for: {failed}")
        return 1
    print(f"All {len(types)} template(s) built and validated successfully.")
    if not args.upload:
        print("Dry run only — pass --upload to publish them to the ORGOS LIBRARY.")
    return 0


if __name__ == "__main__":
    sys.exit(asyncio.run(main()))
