#!/usr/bin/env python3
# =============================================================================
# scripts/build_cdi_reference.py — the document standard, as data
#
# Reads the CDI Procedure (DRG-QI-PRO-CDI) and writes doc_standards/
# cdi_reference.json: the validation checks (V01–V24), required sections,
# Annex A (document register fields), Annex B (reference data) and Annex C
# (control fields). The Document Standards page shows this, and the checks
# read it, so what authors are told is what is enforced — and both come from
# the CDI document, not from anyone's memory of it.
#
# Re-run when the CDI Procedure is reissued:
#   python scripts/build_cdi_reference.py "docs/Control of Documented Information Procedure.docx"
# =============================================================================

import json
import os
import re
import sys

from docx import Document

OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                   "doc_standards", "cdi_reference.json")


def rows(table):
    out = []
    for r in table.rows:
        cells = list(dict.fromkeys(c.text.strip().replace("\n", " ") for c in r.cells))
        if any(cells):
            out.append(cells)
    return out


def table_by_header(doc, *header):
    for t in doc.tables:
        first = [c.text.strip() for c in t.rows[0].cells]
        first = list(dict.fromkeys(first))
        if first[:len(header)] == list(header):
            return rows(t)[1:]
    raise SystemExit(f"table not found: {header}")


def main(path: str) -> None:
    d = Document(path)
    checks = [
        {"id": r[0], "name": r[1], "enforcement": r[2]}
        for r in table_by_header(d, "No.", "Check", "Enforcement") if re.match(r"V\d\d$", r[0])
    ]
    sections = [{"section": r[0], "content": r[1], "required": r[2]}
                for r in table_by_header(d, "Section", "Content", "Required")]
    # Annex A comes first among the "Field | Rule | Mandatory" tables, Annex C last
    field_tables = [t for t in d.tables
                    if list(dict.fromkeys(c.text.strip() for c in t.rows[0].cells))[:3] == ["Field", "Rule", "Mandatory"]]
    register = [{"field": r[0], "rule": r[1], "mandatory": r[2]} for r in rows(field_tables[0])[1:]]
    control = [{"field": r[0], "rule": r[1], "mandatory": r[2]} for r in rows(field_tables[-1])[1:]]
    types = [{"code": r[0], "type": r[1], "layer": r[2]}
             for r in table_by_header(d, "Code", "Type", "Layer or tier")]
    functions = [{"code": r[0], "function": r[1]} for r in table_by_header(d, "Code", "Function")]
    retired = [{"code": r[0], "former": r[1], "current": r[2] if len(r) > 2 else ""}
               for r in table_by_header(d, "Retired code", "Former meaning", "Current code")]
    routes = [{"domain": r[0], "scope": r[1], "first_pass": r[2], "final_approval": r[3]}
              for r in table_by_header(d, "Domain", "Family Policy scope (route follows Primary Domain)",
                                       "First Pass", "Final Approval") if len(r) >= 4]
    data = {
        "source": {"code": "DRG-QI-PRO-CDI", "file": os.path.basename(path)},
        "checks": checks, "required_sections": sections,
        "annex_a_register_fields": register, "annex_c_control_fields": control,
        "annex_b": {"document_types": types, "function_codes": functions,
                    "retired_codes": retired, "domain_routes": routes},
    }
    # A reissued Procedure that changes shape must fail here, not write a hollow file.
    expect = {"checks": (24, 24), "required_sections": (7, 7), "annex_a_register_fields": (20, 40),
              "annex_c_control_fields": (10, 25)}
    for key, (lo, hi) in expect.items():
        n = len(data[key])
        if not lo <= n <= hi:
            raise SystemExit(f"{key}: found {n}, expected {lo}-{hi}. Has the Procedure's layout changed?")
    for key, lo in (("document_types", 8), ("function_codes", 8), ("retired_codes", 5), ("domain_routes", 10)):
        if len(data["annex_b"][key]) < lo:
            raise SystemExit(f"annex_b.{key}: found {len(data['annex_b'][key])}, expected at least {lo}")
    if {c["enforcement"] for c in checks} - {"BLOCK", "FLAG", "REPORT"}:
        raise SystemExit("a check has an unknown enforcement level")
    with open(OUT, "w") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)
    print(f"{len(checks)} checks, {len(sections)} sections, {len(register)} register fields, "
          f"{len(control)} control fields, {len(types)} types, {len(functions)} functions, "
          f"{len(retired)} retired codes, {len(routes)} domain routes → {OUT}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "docs/Control of Documented Information Procedure.docx")
