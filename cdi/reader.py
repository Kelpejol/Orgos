# =============================================================================
# cdi/reader.py — read a controlled document into paragraphs and tables
#
# The parser needs the document's STRUCTURE, not a flat string: a control
# block is a two-column table, a Roles table pairs a role with its duty, an
# Annex B row ties a record to the step that creates it. Flattening loses all
# of that, so this returns the body as an ordered list of Paragraph and Table
# blocks.
#
# Word (.docx)
#   Read straight from the document XML rather than through python-docx's
#   text properties, for one reason that matters: python-docx silently drops
#   text inside tracked changes. A control block row added as a tracked
#   insertion — the CDI Standard's "Child Documents" row is one — comes back
#   EMPTY, with no error. Documents under review are full of tracked changes,
#   so this reads them as if the changes were accepted (insertions kept,
#   deletions dropped) — the same view mammoth gives the CDI checker — and
#   reports how many there are so the parser can say so.
#
# PDF
#   pdfplumber recovers the tables, but a PDF table has no real columns: cells
#   come back padded with empty spacer columns, and a two-line header ("At" /
#   "step") arrives as two rows. Tables are normalised here; because the
#   column POSITIONS are not trustworthy afterwards, they are marked
#   positional=False and the parser aligns them by order and checks the count.
#   Paragraphs are rebuilt from text lines outside the tables, and running
#   page headers and footers are removed.
#
# Pure apart from parsing the bytes it is given. No network, no settings.
# =============================================================================

from __future__ import annotations

import io
import logging
import re
from collections import Counter
from dataclasses import dataclass, field
from typing import Iterator, Union

logger = logging.getLogger(__name__)


@dataclass(frozen=True)
class Paragraph:
    text: str
    style: str = ""          # the Word style name ("Heading 3"); "" for PDF

    @property
    def is_heading_style(self) -> bool:
        s = self.style.lower()
        return s.startswith("heading") or s == "title"


@dataclass(frozen=True)
class Table:
    rows: tuple[tuple[str, ...], ...]
    # True when a cell's index is its real column (Word). False when the
    # table was rebuilt from a PDF and only the ORDER of non-empty cells holds.
    positional: bool = True


Block = Union[Paragraph, Table]


@dataclass
class DocumentContent:
    blocks: list[Block]
    format: str                         # "docx" | "pdf"
    tracked_insertions: int = 0
    tracked_deletions: int = 0
    warnings: list[str] = field(default_factory=list)

    @property
    def tracked_changes(self) -> int:
        return self.tracked_insertions + self.tracked_deletions

    def paragraphs(self) -> Iterator[tuple[int, Paragraph]]:
        for i, b in enumerate(self.blocks):
            if isinstance(b, Paragraph):
                yield i, b

    def tables(self) -> Iterator[tuple[int, Table]]:
        for i, b in enumerate(self.blocks):
            if isinstance(b, Table):
                yield i, b


class UnsupportedDocument(ValueError):
    """The file is not a format the parser can read structurally."""


def read_document(file_bytes: bytes, filename: str) -> DocumentContent:
    ext = (filename or "").lower().rsplit(".", 1)[-1]
    if ext == "docx":
        return _read_docx(file_bytes)
    if ext == "pdf":
        return _read_pdf(file_bytes)
    if ext == "doc":
        raise UnsupportedDocument(
            "A legacy .doc file has no readable structure. Save it as .docx."
        )
    raise UnsupportedDocument(f"'.{ext}' files are not supported. Use .docx or .pdf.")


# =============================================================================
#  Word
# =============================================================================

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _q(tag: str) -> str:
    return f"{{{_W}}}{tag}"


_T, _TAB, _BR, _CR, _NBH = _q("t"), _q("tab"), _q("br"), _q("cr"), _q("noBreakHyphen")
_P, _TBL, _TR, _TC, _SDT, _SDT_CONTENT = (
    _q("p"), _q("tbl"), _q("tr"), _q("tc"), _q("sdt"), _q("sdtContent"),
)
# Text inside these is pending REMOVAL; the accepted view drops it.
_REMOVED = {_q("del"), _q("moveFrom")}


def _removed(node, stop) -> bool:
    """True if `node` sits inside a tracked deletion (below `stop`)."""
    parent = node.getparent()
    while parent is not None and parent is not stop:
        if parent.tag in _REMOVED:
            return True
        parent = parent.getparent()
    return False


def _accepted_text(el) -> str:
    """The text of a Word element as it reads with tracked changes accepted."""
    parts: list[str] = []
    for node in el.iter(_T, _TAB, _BR, _CR, _NBH):
        if _removed(node, el):
            continue
        if node.tag == _T:
            parts.append(node.text or "")
        elif node.tag == _TAB:
            parts.append("\t")
        elif node.tag == _NBH:
            parts.append("-")
        else:
            parts.append("\n")
    return "".join(parts)


def _clean(text: str) -> str:
    """Collapse runs of whitespace inside a line, keep line breaks."""
    lines = [re.sub(r"[ \t ]+", " ", ln).strip() for ln in text.split("\n")]
    return "\n".join(ln for ln in lines if ln)


def _read_docx(file_bytes: bytes) -> DocumentContent:
    from docx import Document

    try:
        doc = Document(io.BytesIO(file_bytes))
    except Exception as exc:
        raise UnsupportedDocument(f"This file could not be opened as Word: {exc}") from exc

    style_names = {s.style_id: s.name for s in doc.styles}
    body = doc.element.body
    blocks: list[Block] = []

    def para(p) -> Paragraph:
        style = ""
        ppr = p.find(_q("pPr"))
        if ppr is not None:
            ps = ppr.find(_q("pStyle"))
            if ps is not None:
                style = style_names.get(ps.get(_q("val")), ps.get(_q("val")) or "")
        return Paragraph(_clean(_accepted_text(p)), style)

    def table(tbl) -> Table:
        rows: list[tuple[str, ...]] = []
        for tr in tbl.iter(_TR):
            if tr.getparent() is not tbl and tr.getparent().tag != _SDT_CONTENT:
                continue  # a row of a nested table — read as part of its cell
            trpr = tr.find(_q("trPr"))
            if trpr is not None and trpr.find(_q("del")) is not None:
                continue  # the whole row is a tracked deletion
            cells = [
                _clean("\n".join(_accepted_text(p) for p in tc.iter(_P)))
                for tc in tr.iter(_TC)
                if tc.getparent() is tr or tc.getparent().tag == _SDT_CONTENT
            ]
            if any(cells):
                rows.append(tuple(cells))
        return Table(tuple(rows), positional=True)

    def walk(parent) -> None:
        for child in parent.iterchildren():
            if child.tag == _P:
                p = para(child)
                # A table-of-contents entry repeats a heading with a page
                # number; read as content, "4.0 Title (BCDR-01)" would count twice.
                if p.text and not p.style.lower().startswith("toc"):
                    blocks.append(p)
            elif child.tag == _TBL:
                t = table(child)
                if t.rows:
                    blocks.append(t)
            elif child.tag == _SDT:
                content = child.find(_SDT_CONTENT)
                if content is not None:
                    walk(content)

    walk(body)
    return DocumentContent(
        blocks=blocks,
        format="docx",
        tracked_insertions=sum(1 for _ in body.iter(_q("ins"), _q("moveTo"))),
        tracked_deletions=sum(1 for _ in body.iter(_q("del"), _q("moveFrom"))),
    )


# =============================================================================
#  PDF
# =============================================================================

# A line that begins a new paragraph rather than continuing the last one.
_NEW_PARAGRAPH = re.compile(
    r"""^(?:
        \d+\.\d+(?:\.\d+)*\s           # numbered clause or heading: 4.0, 4.1, 5.3.2
      | Annex\s+[A-Z]\b                # Annex A. Control map
      | [A-Z][A-Z0-9]{1,7}-\d{2,3}\s   # a control heading: CDI-001 Title
      | Risk\ statement\b
      | [A-Z][A-Za-z /&()-]{1,40}:\s   # a labelled field: Requirement: …
    )""",
    re.VERBOSE,
)
_PAGE_NUMBER = re.compile(r"^Page\s+\d+\s+of\s+\d+$", re.IGNORECASE)
# A PDF contents line: "4.0 Business impact assessment ........ 3"
_CONTENTS_LINE = re.compile(r"(?:\.\s?){4,}\s*\d+$|…\s*\d+$")


def normalise_pdf_rows(raw_rows: list[list[str]]) -> tuple[tuple[str, ...], ...]:
    """
    Turn a PDF table's raw rows into rows of their non-empty cells, in order.

    1. A header split over two lines ("At" then "step", "Approved" then "By")
       is merged first, while positions still line up with the header.
    2. Empty spacer cells are then dropped from every row.

    After this only the ORDER of cells is meaningful.
    """
    rows = [[re.sub(r"\s+", " ", (c or "")).strip() for c in r] for r in raw_rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return ()

    header = list(rows[0])
    header_cols = [j for j, c in enumerate(header) if c]
    i = 1
    while i < len(rows):
        filled = [(j, c) for j, c in enumerate(rows[i]) if c]
        # A split header word sits under SOME header cells, never all of them,
        # and never under the first: a label/value table's second row fills
        # the same columns as its first, and is data.
        is_continuation = (
            0 < len(filled) < len(header_cols)
            and header_cols[0] not in {j for j, _ in filled}
            and all(
                j < len(header) and header[j]
                and len(c.split()) <= 2
                and not re.search(r"\d|REC-|DRG-", c)
                for j, c in filled
            )
        )
        if not is_continuation:
            break
        for j, c in filled:
            header[j] = f"{header[j]} {c}"
        i += 1

    return tuple(
        tuple(c for c in r if c) for r in [header] + rows[i:]
    )


def _join_lines(lines: list[str]) -> list[str]:
    """Rebuild paragraphs from wrapped PDF lines."""
    paragraphs: list[str] = []
    for line in lines:
        if not paragraphs or _NEW_PARAGRAPH.match(line):
            paragraphs.append(line)
        elif paragraphs[-1].endswith("-") and (line[:1].islower() or line[:1].isdigit()):
            # A word or an ID broken at its hyphen: "Contract-" / "owning",
            # "REC-" / "0678". Rejoined without a space, or the ID is lost.
            paragraphs[-1] = paragraphs[-1] + line
        else:
            paragraphs[-1] = f"{paragraphs[-1]} {line}"
    return paragraphs


def _running_lines(pages: list[list[str]]) -> set[str]:
    """Lines repeated at the top or bottom of several pages — headers/footers."""
    edges = Counter()
    for lines in pages:
        for ln in set(lines[:2] + lines[-2:]):
            edges[ln] += 1
    return {ln for ln, n in edges.items() if n >= 2 and len(pages) >= 2}


def _read_pdf(file_bytes: bytes) -> DocumentContent:
    import pdfplumber

    try:
        pdf = pdfplumber.open(io.BytesIO(file_bytes))
    except Exception as exc:
        raise UnsupportedDocument(f"This file could not be opened as PDF: {exc}") from exc

    # Per page: text lines outside tables (with their vertical position) and
    # tables (with theirs), so the two can be put back in reading order.
    page_lines: list[list[tuple[float, str]]] = []
    page_tables: list[list[tuple[float, Table]]] = []
    with pdf:
        for page in pdf.pages:
            found = page.find_tables()
            boxes = [t.bbox for t in found]

            def outside(obj, boxes=boxes) -> bool:
                cx = (obj.get("x0", 0) + obj.get("x1", 0)) / 2
                cy = (obj.get("top", 0) + obj.get("bottom", 0)) / 2
                return not any(x0 <= cx <= x1 and top <= cy <= bottom
                               for x0, top, x1, bottom in boxes)

            lines = [
                (ln["top"], re.sub(r"\s+", " ", ln["text"]).strip())
                for ln in page.filter(outside).extract_text_lines()
            ]
            page_lines.append([(y, t) for y, t in lines if t])
            page_tables.append([
                (t.bbox[1], Table(normalise_pdf_rows(t.extract()), positional=False))
                for t in found
            ])

    running = _running_lines([[t for _, t in lines] for lines in page_lines])

    blocks: list[Block] = []
    for lines, tables in zip(page_lines, page_tables):
        kept = [(y, t) for y, t in lines
                if not _PAGE_NUMBER.match(t) and not _CONTENTS_LINE.search(t)
                and t not in running]
        items: list[tuple[float, Union[str, Table]]] = [*kept, *tables]
        items.sort(key=lambda it: it[0])
        pending: list[str] = []
        for _, item in items:
            if isinstance(item, Table):
                blocks.extend(Paragraph(p) for p in _join_lines(pending))
                pending = []
                if item.rows:
                    blocks.append(item)
            else:
                pending.append(item)
        blocks.extend(Paragraph(p) for p in _join_lines(pending))

    # A paragraph broken across a page boundary arrives as two; rejoin it.
    merged: list[Block] = []
    for b in blocks:
        if (merged and isinstance(b, Paragraph) and isinstance(merged[-1], Paragraph)
                and not _NEW_PARAGRAPH.match(b.text)
                and not re.search(r"[.:;)]$", merged[-1].text)):
            merged[-1] = Paragraph(f"{merged[-1].text} {b.text}")
        else:
            merged.append(b)

    return DocumentContent(blocks=merged, format="pdf")
