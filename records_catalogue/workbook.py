# =============================================================================
# records_catalogue/workbook.py — read the Enterprise Records Catalogue workbook
#
# The Catalogue is maintained as an Excel workbook; its master sheet is
# "Enterprise Records Catalogue", with a two-line banner above the header row.
# This reads that sheet into RecordEntry objects. The derived sheets
# (Findability View, Retention & Disposition, …) are views over the same Record
# IDs (Catalogue rule GR-10), so they are not read separately.
#
# Pure: bytes in, entries out.
# =============================================================================

from __future__ import annotations

import io
from typing import Any

from records_catalogue.model import RecordEntry

MASTER_SHEET = "Enterprise Records Catalogue"

# The workbook's column headers, exactly as written, → RecordEntry fields.
COLUMNS: dict[str, str] = {
    "Proposed Record ID": "record_id",
    "Canonical Record Name": "name",
    "Aliases / Source Names": "aliases",
    "Controlled Record Type": "record_type",
    "Type Mapping Confidence": "type_mapping_confidence",
    "Admission Status": "admission_status",
    "Discovery Proposed Type": "discovery_proposed_type",
    "Source Count": "source_count",
    "Source Documents": "source_documents",
    "Source / Process References": "source_process_references",
    "Source URLs": "source_urls",
    "Privacy Overlay": "privacy_overlay",
    "Candidate Findability Statement": "findability_statement",
    "Findability Validation": "findability_validation",
    "Candidate Retention Statement": "retention_statement",
    "Retention Validation": "retention_validation",
    "Process Custodian": "custodian",
    "Authoritative Location / System": "location",
    "Retrieval Instruction": "retrieval_instruction",
    "Retention Period / Rule": "retention_rule",
    "Retention Trigger": "retention_trigger",
    "Retention Authority": "retention_authority",
    "Disposition": "disposition",
    "Destruction Method": "destruction_method",
    "Destruction Evidence": "destruction_evidence",
    "Legal Hold / Suspension": "legal_hold",
    "Specialist Register Reference": "specialist_register_reference",
    "Mapping Rationale": "mapping_rationale",
    "Record Validation Status": "record_validation_status",
}


class WorkbookError(ValueError):
    """The workbook is not the Catalogue, or its layout has changed."""


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def read_catalogue(data: bytes) -> list[RecordEntry]:
    """
    Every record on the master sheet. Raises WorkbookError if the sheet or a
    column is missing — a changed layout must stop an import, not half-load it.
    """
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise WorkbookError(f"Not a readable Excel workbook: {exc}") from exc
    if MASTER_SHEET not in wb.sheetnames:
        raise WorkbookError(f"No '{MASTER_SHEET}' sheet; found {wb.sheetnames}.")

    rows = list(wb[MASTER_SHEET].iter_rows(values_only=True))
    header_at = next(
        (i for i, r in enumerate(rows[:15]) if r and _text(r[0]) == "Proposed Record ID"), None
    )
    if header_at is None:
        raise WorkbookError("Could not find the header row ('Proposed Record ID').")

    header = [_text(h) for h in rows[header_at]]
    missing = [c for c in COLUMNS if c not in header]
    if missing:
        raise WorkbookError(f"The Catalogue is missing column(s): {', '.join(missing)}.")
    index = {COLUMNS[h]: i for i, h in enumerate(header) if h in COLUMNS}

    entries: list[RecordEntry] = []
    for row in rows[header_at + 1:]:
        if not row or not _text(row[index["record_id"]]).startswith("REC-"):
            continue
        values = {f: _text(row[i]) if i < len(row) else "" for f, i in index.items()}
        values["aliases"] = [a.strip() for a in values["aliases"].split("|") if a.strip()]
        entries.append(RecordEntry(**values))
    return entries
