# =============================================================================
# tests/doc_standards/test_role_fix.py — "use this role" edits a controlled document
#
# It must change a role FIELD and nothing else: not the same words in running
# text, not other roles, not formatting. And it must refuse rather than guess.
# =============================================================================

import io

import pytest
from docx import Document

from doc_standards.role_fix import RoleFixError, replace_role
from tests.cdi.builders import DocBuilder

OLD, NEW = "Info Sec Officer", "Information Security Officer"


def texts(data: bytes) -> list[str]:
    doc = Document(io.BytesIO(data))
    out = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            out += [c.text for c in r.cells]
    return out


def control_doc(owner_line=f"Control owner: {OLD}.", extra=()):
    b = DocBuilder()
    b.control_block(title="Test Standard", code="DRG-ISMS-STD-TST", layer="Layer 3 Standard", domain="d",
                    owner=OLD, first_pass=OLD, final_approval="Chief Compliance Officer", classification="c")
    b.control("TST-01", "One", requirement="The Info Sec Officer shall review it.", trigger="Creation.")
    b.para(owner_line)
    for line in extra:
        b.para(line)
    return b.bytes()


def test_a_control_owner_line_is_changed():
    fix = replace_role(control_doc(), OLD, NEW)
    assert f"Control owner: {NEW}." in texts(fix.data) and fix.count >= 1


def test_the_same_words_in_running_text_are_left_alone():
    fix = replace_role(control_doc(), OLD, NEW)
    requirement = next(t for t in texts(fix.data) if t.startswith("Requirement:"))
    assert requirement == "Requirement: The Info Sec Officer shall review it."          # untouched


def test_a_header_table_owner_and_approver_rows_are_changed_but_other_rows_are_not():
    fix = replace_role(control_doc(), OLD, NEW)
    cells = texts(fix.data)
    assert cells.count(NEW) >= 2                       # Document owner and First pass approval
    assert "Chief Compliance Officer" in cells         # another role, untouched
    assert "Test Standard" in cells


def test_the_older_accountable_executing_line_changes_only_the_accountable_part():
    fix = replace_role(control_doc(owner_line=f"Accountable / Executing: {OLD}; IT and application owners."), OLD, NEW)
    assert f"Accountable / Executing: {NEW}; IT and application owners." in texts(fix.data)


def test_several_labelled_segments_in_one_paragraph_change_only_the_role_one():
    line = f"Governing authority: DRG-ISMS-POL-ACCESS. Control owner: {OLD}. Evidence: {OLD} log."
    fix = replace_role(control_doc(owner_line=line), OLD, NEW)
    assert f"Governing authority: DRG-ISMS-POL-ACCESS. Control owner: {NEW}. Evidence: {OLD} log." in texts(fix.data)


def test_a_role_split_across_formatting_runs_is_still_changed_and_keeps_its_formatting():
    b = DocBuilder()
    b.control_block(title="T", code="DRG-ISMS-STD-TST", layer="Layer 3", domain="d", owner="X", classification="c")
    p = b.doc.add_paragraph()
    p.add_run("Control owner: ")
    bold = p.add_run("Info Sec "); bold.bold = True
    p.add_run("Officer.")
    fix = replace_role(b.bytes(), OLD, NEW)
    changed = next(p for p in Document(io.BytesIO(fix.data)).paragraphs if p.text.startswith("Control owner:"))
    assert changed.text == f"Control owner: {NEW}."
    assert any(r.bold for r in changed.runs)            # the run the role began in kept its formatting


def test_a_control_table_role_column_is_changed():
    b = DocBuilder()
    b.control_block(title="T", code="DRG-ISMS-STD-TST", layer="Layer 3", domain="d", owner="X", classification="c")
    b.table([["Control ID", "Requirement", "Control owner"], ["TST-01", f"The {OLD} shall.", OLD], ["TST-02", "Other.", "HR Manager"]])
    fix = replace_role(b.bytes(), OLD, NEW)
    cells = texts(fix.data)
    assert NEW in cells and "HR Manager" in cells and f"The {OLD} shall." in cells    # prose cell untouched


def test_a_roles_table_is_changed():
    b = DocBuilder()
    b.control_block(title="T", code="DRG-ISMS-STD-TST", layer="Layer 3", domain="d", owner="X", classification="c")
    b.table([["Role", "Responsibility under this Procedure"], [OLD, f"Reviews. The {OLD} approves."]])
    fix = replace_role(b.bytes(), OLD, NEW)
    cells = texts(fix.data)
    assert NEW in cells and f"Reviews. The {OLD} approves." in cells


def test_a_document_that_names_the_role_only_in_prose_is_not_changed_and_says_so():
    b = DocBuilder()
    b.control_block(title="T", code="DRG-ISMS-STD-TST", layer="Layer 3", domain="d", owner="X", classification="c")
    b.para(f"The {OLD} approves access requests.")
    data = b.bytes()
    fix = replace_role(data, OLD, NEW)
    assert fix.count == 0 and fix.data == data


def test_only_the_named_role_is_changed():
    fix = replace_role(control_doc(extra=["Control owner: HR Manager."]), OLD, NEW)
    assert "Control owner: HR Manager." in texts(fix.data)


def test_nothing_is_added_or_removed():
    original = control_doc()
    fix = replace_role(original, OLD, NEW)
    a, b = Document(io.BytesIO(original)), Document(io.BytesIO(fix.data))
    assert (len(a.paragraphs), len(a.tables)) == (len(b.paragraphs), len(b.tables))


def test_it_refuses_empty_or_unchanged_input():
    for a, b in (("", NEW), (OLD, ""), (OLD, OLD)):
        with pytest.raises(RoleFixError):
            replace_role(control_doc(), a, b)


def test_a_change_that_would_alter_more_than_the_role_is_refused(monkeypatch):
    import doc_standards.role_fix as rf
    real = rf._fix_tables

    def vandal(doc, role, new, fix):
        real(doc, role, new, fix)
        doc.paragraphs[0].runs[0].text = "TAMPERED"
        fix.count += 1
    monkeypatch.setattr(rf, "_fix_tables", vandal)
    with pytest.raises(RoleFixError, match="more than the role"):
        replace_role(control_doc(), OLD, NEW)
