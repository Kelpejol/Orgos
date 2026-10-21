# Shared builders for the document-standards tests: a master template whose cover
# reads like the real one, and v06-shaped documents with a hand-typed header block.

import io

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from tests.cdi.builders import DocBuilder


def master_template() -> bytes:
    """A cover with the real template's labels and markers, then the body placeholder."""
    doc = Document()
    doc.add_paragraph("DRAGNET SOLUTIONS")
    doc.add_paragraph("{{ domain }} · {{ doc_type }}")
    doc.add_paragraph("{{ title }}")
    rows = [("Document Title", "{{ title }}"), ("Document Code", "{{ doc_code }}"),
            ("Document Type / Layer", "{{ type_layer }}"), ("Domain", "{{ domain }}"),
            ("Parent Document", "{{ parent_document }}"), ("Classification", "{{ classification }}"),
            ("Distribution", "{{ distribution }}"), ("Owner", "{{ owner }}"),
            ("First Pass Approval", "{{ first_pass_approval }}"), ("Final Approval", "{{ final_approval }}"),
            ("Version", "{{ version }}"), ("Effective Date", "{{ effective_date }}"),
            ("Next Review Due", "{{ next_review_due }}")]
    t = doc.add_table(rows=len(rows) + 3, cols=2)
    for i, (a, b) in enumerate(rows):
        t.cell(i, 0).text, t.cell(i, 1).text = a, b
    t.cell(len(rows), 0).text = "{%tr for s in standards %}"
    t.cell(len(rows) + 1, 0).text, t.cell(len(rows) + 1, 1).text = "{{ s.name }}", "{{ s.reference }}"
    t.cell(len(rows) + 2, 0).text = "{%tr endfor %}"
    doc.add_paragraph("REVISION HISTORY")
    h = doc.add_table(rows=4, cols=4)
    for j, text in enumerate(("Version", "Date", "Purpose of Change", "Approved By")):
        h.cell(0, j).text = text
    h.cell(1, 0).text = "{%tr for r in revision_history %}"
    for j, text in enumerate(("{{ r.version }}", "{{ r.date }}", "{{ r.purpose }}", "{{ r.approved_by }}")):
        h.cell(2, j).text = text
    h.cell(3, 0).text = "{%tr endfor %}"
    doc.add_paragraph("[Body content begins here — 1.0 Purpose & Scope, etc.]")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def v06_document(code="DRG-ISMS-STD-TST", title="Test Standard", *, ids=("TST-01", "TST-02"),
                 version="v1.0", implemented=None, mapping="ISO 27001 A.5.16", bookmark=False,
                 extra_body=None) -> bytes:
    """A document the way v06 authors write them: a title, a header-block table, then the body."""
    b = DocBuilder()
    b.heading(title, level=1)
    fields = dict(title=title, code=code, layer="Layer 3 Standard", domain="Information Security",
                  parent="Access Control Policy (DRG-ISMS-POL-ACCESS)", owner="Information Security Officer",
                  classification="Internal Use Only", status="Working draft",
                  standards="ISO/IEC 27001:2022 Annex A 5.16", version=version)
    if implemented:
        fields["implemented"] = implemented
    b.control_block(**fields)
    b.heading("Scope and authority", level=1)
    b.para("This standard defines one canonical control statement per Control ID.")
    b.heading("Controls", level=1)
    for cid in ids:
        kw = dict(requirement=f"The Information Security Officer shall maintain {cid}.",
                  trigger="Creation.", responsibility="Information Security Officer; IT owners.",
                  evidence="Inventory entry. Record IDs: REC-0691.", risk="An account could remain unidentified.",
                  failure_route="Quarantine the identity.", exception_route="Recorded in REC-0696.")
        if mapping:
            kw["mapping"] = mapping
        b.control(cid, f"Control {cid}", **kw)
    for text in extra_body or []:
        b.para(text)
    if bookmark:
        body = b.doc.element.body
        body.insert(0, parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="31" w:name="section-top"/>'))
        body.insert(len(body) - 1, parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="31"/>'))
    return b.bytes()
