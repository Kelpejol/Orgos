# =============================================================================
# tests/agents/policy_drafter/test_v06_drafting.py — AI drafter on CDT v06
#
# Proves the rebuilt drafter: document codes drop the trailing serial/year
# (v06 §4.0), a genuine code collision gets a disambiguator rather than
# reviving the old numbering scheme, and draft_document() builds the .docx on
# top of the document type's master CDT template (fetched + validated),
# seeding sensible initial cover facts — never a hand-rolled, drifting cover.
# =============================================================================

import io
from unittest.mock import AsyncMock, patch

import pytest
from docx import Document

from agents.policy_drafter.service import (
    _generate_doc_code,
    draft_document,
    generate_doc_code_base,
)
from agents.policy_drafter import router as drafter_router
from lifecycle.merge import scan_markers


def _fixture_master_template() -> bytes:
    """A minimal v06-shaped master (cover w/ standards tail + history), reusing
    the same 3-row docxtpl loop pattern proven in tests/lifecycle/test_merge.py."""
    doc = Document()
    ctbl = doc.add_table(rows=0, cols=2)
    for label, marker in [
        ("Document Title", "{{ title }}"), ("Document Code", "{{ doc_code }}"),
        ("Version", "{{ version }}"), ("Classification", "{{ classification }}"),
    ]:
        row = ctbl.add_row()
        row.cells[0].text = label
        row.cells[1].text = marker
    for tag, cells in [
        ("{%tr for s in standards %}", None), (None, ["{{ s.name }}", "{{ s.reference }}"]),
        ("{%tr endfor %}", None),
    ]:
        row = ctbl.add_row()
        if tag:
            row.cells[0].text = tag
        else:
            row.cells[0].text, row.cells[1].text = cells
    htbl = doc.add_table(rows=0, cols=4)
    for tag_or_cells in [
        ("{%tr for r in revision_history %}",), None, ("{%tr endfor %}",)
    ]:
        row = htbl.add_row()
        if tag_or_cells:
            row.cells[0].text = tag_or_cells[0]
        else:
            for i, f in enumerate(["version", "date", "purpose", "approved_by"]):
                row.cells[i].text = "{{ r.%s }}" % f
    doc.add_paragraph("BODY GOES HERE")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# -----------------------------------------------------------------------------
#  Document code — v06 format, no trailing serial/year
# -----------------------------------------------------------------------------

def test_generate_doc_code_has_no_suffix():
    code = _generate_doc_code("Quality & Compliance", "Procedure", "Control of Documented Information")
    assert code.startswith("DRG-")
    parts = code.split("-")
    assert not (parts[-1].isdigit() and len(parts[-1]) == 2), \
        f"v06 codes must not end in a bare 2-digit serial/year: {code}"


def test_generate_doc_code_disambiguator_appends_to_id():
    base = generate_doc_code_base("Quality & Compliance", "Procedure", "Control of Documented Information")
    with_disambiguator = _generate_doc_code(
        "Quality & Compliance", "Procedure", "Control of Documented Information", "2",
    )
    assert with_disambiguator == base + "2"


def test_combined_doc_type_gets_compound_pol_pro_code():
    """
    A "Combined" (policy + procedure in one) document gets the real
    POL-PRO compound code — the legacy naming convention
    agents.cdi_checker.service.is_valid_doc_code()/is_combined_document()
    and the extractor's CODE_PREFIX_MAP already recognise — not a plain
    POL code indistinguishable from an ordinary Policy.
    """
    from agents.cdi_checker.service import is_combined_document, is_valid_doc_code

    code = generate_doc_code_base("QI", "Combined", "New Control Access Policy")
    assert "-POL-PRO-" in code
    assert is_valid_doc_code(code)
    assert is_combined_document("", code)


# -----------------------------------------------------------------------------
#  Collision detection (_next_disambiguator) — v06 has no serial/year fallback
# -----------------------------------------------------------------------------

async def test_no_disambiguator_needed_when_code_is_free():
    with patch("agents.policy_drafter.router.get_list_items", new_callable=AsyncMock) as mock_items:
        mock_items.return_value = [{"fields": {"DocumentCode": "DRG-QI-PRO-OTHER"}}]
        result = await drafter_router._next_disambiguator("DRG-QI-PRO-CDI")
    assert result == ""


async def test_disambiguator_increments_past_existing_collisions():
    with patch("agents.policy_drafter.router.get_list_items", new_callable=AsyncMock) as mock_items:
        mock_items.return_value = [
            {"fields": {"DocumentCode": "DRG-QI-PRO-CDI"}},
            {"fields": {"DocumentCode": "DRG-QI-PRO-CDI2"}},
        ]
        result = await drafter_router._next_disambiguator("DRG-QI-PRO-CDI")
    assert result == "3"  # "" and "2" both taken


async def test_disambiguator_defaults_empty_on_query_failure():
    with patch("agents.policy_drafter.router.get_list_items", new_callable=AsyncMock) as mock_items:
        mock_items.side_effect = Exception("Graph down")
        result = await drafter_router._next_disambiguator("DRG-QI-PRO-CDI")
    assert result == ""


# -----------------------------------------------------------------------------
#  draft_document() — builds on the master template, seeds cover facts
# -----------------------------------------------------------------------------

async def test_draft_document_builds_on_master_template_with_seeded_cover():
    master = _fixture_master_template()
    with patch("agents.policy_drafter.service.load_master_template", new_callable=AsyncMock) as mock_load, \
         patch("agents.policy_drafter.service.llm_generate", new_callable=AsyncMock) as mock_llm:
        mock_load.return_value = master
        mock_llm.return_value = "Generated section text."

        draft = await draft_document(
            title="Control of Documented Information Procedure",
            doc_type="Procedure",
            department="Quality & Compliance",
            standards_mapping="ISO 9001:2015, ISO 27001:2022, NDPA 2023",
        )

    mock_load.assert_awaited_once_with("Procedure")

    # Seeded cover facts — a starting point, not a final answer.
    assert draft["cdt_cover"]["version"] == "Draft"
    assert draft["cdt_cover"]["domain"] == "Quality & Compliance"
    assert draft["cdt_cover"]["classification"] == "Internal use only"
    assert [s["name"] for s in draft["cdt_cover"]["standards"]] == [
        "ISO 9001:2015", "ISO 27001:2022", "NDPA 2023",
    ]
    assert draft["revision_history"][0]["version"] == "Draft"
    assert draft["revision_history"][0]["purpose"] == "Initial AI-generated draft"

    # The built .docx is the MASTER TEMPLATE plus body content — cover markers
    # are still LIVE (unfilled). The AI never touches them.
    docx_bytes = draft["docx_buffer"].getvalue()
    markers = scan_markers(docx_bytes)
    assert any("doc_code" in m for m in markers)
    assert any("standards" in m for m in markers)
    d = Document(io.BytesIO(docx_bytes))
    body_text = " ".join(p.text for p in d.paragraphs)
    assert "BODY GOES HERE" in body_text          # master's own content preserved
    assert "1. PURPOSE" in body_text.upper()      # AI body appended after it
    assert "Generated section text." in body_text


async def test_draft_document_propagates_template_not_found():
    from lifecycle.templates import TemplateNotFoundError
    with patch("agents.policy_drafter.service.load_master_template", new_callable=AsyncMock) as mock_load:
        mock_load.side_effect = TemplateNotFoundError("no master for 'Weird'")
        with pytest.raises(TemplateNotFoundError):
            await draft_document(title="X", doc_type="Weird", department="Y")
