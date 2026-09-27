# =============================================================================
# tests/lifecycle/test_merge.py — CDT merge engine unit tests (v06)
#
# The merge engine is pure (no Graph/SharePoint), so these are plain sync tests
# over in-memory .docx bytes. They prove the compliance-critical guarantees:
# determinism, fail-loud on unresolved/unknown markers, correct row expansion
# (both the revision-history AND standards repeating blocks), XML-safe
# escaping, and Unicode fidelity.
#
# v06: there is no header/footer zone any more (confirmed empty by the
# compliance officer) — CDT is two zones, both in the body: the cover table
# (with the growing `standards` block as its tail) and the revision-history
# table.
# =============================================================================

import hashlib
import io

import pytest
from docx import Document

from lifecycle.merge import (
    MergeError,
    MergeResult,
    MissingContextError,
    TemplateRenderError,
    UnresolvedMarkerError,
    merge,
    render_document,
    scan_markers,
)
from lifecycle.schemas import MergeContext, RevisionHistoryEntry, StandardEntry


# -----------------------------------------------------------------------------
#  Template builders (mimic a hand-authored v06 master: cover + history)
# -----------------------------------------------------------------------------

def _add_loop(table, *, loop_var: str, iterable: str, cols: list[str]):
    """
    Build the 3-row docxtpl row-loop pattern (dedicated for-row / data-row /
    endfor-row) — the ONLY layout under which docxtpl actually repeats a row.
    Mixing the {%tr for %} tag into the same row as data markers produces
    ZERO output rows (verified while rebuilding this file for v06).
    """
    for_row = table.add_row()
    for_row.cells[0].text = f"{{%tr for {loop_var} in {iterable} %}}"
    data_row = table.add_row()
    for i, field in enumerate(cols):
        data_row.cells[i].text = f"{{{{ {loop_var}.{field} }}}}"
    end_row = table.add_row()
    end_row.cells[0].text = "{%tr endfor %}"


def _build_master(extra_body: str | None = None) -> bytes:
    """A minimal v06 master: cover table (+ standards tail) and history table."""
    doc = Document()

    # Cover control table
    doc.add_paragraph("DOCUMENT CONTROL")
    ctbl = doc.add_table(rows=0, cols=2)
    for label, marker in [
        ("Document Title",         "{{ title }}"),
        ("Document Code",          "{{ doc_code }}"),
        ("Document Type / Layer",  "{{ type_layer }}"),
        ("Domain",                 "{{ domain }}"),
        ("Parent Document",        "{{ parent_document }}"),
        ("Classification",         "{{ classification }}"),
        ("Distribution",           "{{ distribution }}"),
        ("Owner",                  "{{ owner }}"),
        ("First Pass Approval",    "{{ first_pass_approval }}"),
        ("Final Approval",         "{{ final_approval }}"),
        ("Version",                "{{ version }}"),
        ("Effective Date",         "{{ effective_date }}"),
        ("Next Review Due",        "{{ next_review_due }}"),
    ]:
        row = ctbl.add_row()
        row.cells[0].text = label
        row.cells[1].text = marker
    _add_loop(ctbl, loop_var="s", iterable="standards", cols=["name", "reference"])

    # Revision-history table
    doc.add_paragraph("REVISION HISTORY")
    htbl = doc.add_table(rows=1, cols=4)
    for i, h in enumerate(["Version", "Date", "Purpose of Change", "Approved By"]):
        htbl.rows[0].cells[i].text = h
    _add_loop(htbl, loop_var="r", iterable="revision_history",
              cols=["version", "date", "purpose", "approved_by"])

    # Body — no markers
    doc.add_paragraph("1.0 PURPOSE. Body text; contains no markers.")
    if extra_body:
        doc.add_paragraph(extra_body)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _sample_context(**overrides) -> MergeContext:
    base = dict(
        doc_code="DRG-QI-PRO-CDI",
        title="Control of Documented Information Procedure",
        doc_type="Procedure",
        type_layer="Procedure (Layer 4)",
        domain="Quality & Compliance",
        parent_document="Control of Documented Information Policy (Layer 2)",
        classification="Internal use only",
        distribution="All staff",
        owner="Compliance Senior Executive",
        first_pass_approval="Compliance Senior Executive",
        final_approval="Chief Compliance Officer",
        version="06",
        effective_date="August 2026",
        next_review_due="August 2027",
        standards=[
            StandardEntry(name="ISO 9001:2015", reference="Clause 7.5"),
            StandardEntry(name="ISO 27001:2022", reference="Clause 7.5"),
            StandardEntry(name="NDPA 2023", reference="Section 24"),
        ],
        revision_history=[
            RevisionHistoryEntry(version="05", date="June 2026", purpose="Full rewrite",
                                 approved_by="Chief Compliance Officer"),
            RevisionHistoryEntry(version="06", date="August 2026", purpose="Approval authority moved",
                                 approved_by="Chief Compliance Officer"),
        ],
    )
    base.update(overrides)
    return MergeContext(**base)


def _all_text(docx_bytes: bytes) -> str:
    """Concatenate all readable text (body + tables) for assertions."""
    d = Document(io.BytesIO(docx_bytes))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    return " ".join(parts)


# -----------------------------------------------------------------------------
#  Happy path
# -----------------------------------------------------------------------------

def test_merge_fills_cover_and_leaves_no_markers():
    result = merge(_build_master(), _sample_context())
    assert isinstance(result, MergeResult)
    assert scan_markers(result.document) == []          # nothing left over
    text = _all_text(result.document)
    assert "DRG-QI-PRO-CDI" in text
    assert "Chief Compliance Officer" in text
    assert "Internal use only" in text
    assert result.doc_code == "DRG-QI-PRO-CDI"
    assert result.sha256 == hashlib.sha256(result.document).hexdigest()


def test_body_untouched():
    result = merge(_build_master(), _sample_context())
    d = Document(io.BytesIO(result.document))
    assert "1.0 PURPOSE" in " ".join(p.text for p in d.paragraphs)


# -----------------------------------------------------------------------------
#  Determinism (compliance-critical — ME-06)
# -----------------------------------------------------------------------------

def test_merge_is_deterministic_byte_for_byte():
    src = _build_master()
    ctx = _sample_context()
    r1 = merge(src, ctx)
    r2 = merge(src, ctx)
    assert r1.sha256 == r2.sha256
    assert r1.document == r2.document


# -----------------------------------------------------------------------------
#  Repeating-row expansion — revision history AND standards
# -----------------------------------------------------------------------------

def test_history_expands_one_row_per_entry():
    result = merge(_build_master(), _sample_context())
    d = Document(io.BytesIO(result.document))
    history = d.tables[1]  # tables: [0]=cover, [1]=history
    assert len(history.rows) - 1 == 2  # minus header row
    assert history.rows[1].cells[0].text == "05"
    assert history.rows[2].cells[0].text == "06"


def test_empty_history_produces_no_data_rows_without_error():
    result = merge(_build_master(), _sample_context(revision_history=[]))
    d = Document(io.BytesIO(result.document))
    history = d.tables[1]
    assert len(history.rows) == 1  # header row only
    assert scan_markers(result.document) == []


def test_standards_expand_one_row_per_entry():
    """The v06 growing-standards block — replaces the old fixed ISO cells."""
    result = merge(_build_master(), _sample_context())
    d = Document(io.BytesIO(result.document))
    cover = d.tables[0]
    # 13 fixed rows + 3 standards rows
    assert len(cover.rows) == 16
    tail = [tuple(c.text for c in row.cells) for row in cover.rows[13:]]
    assert tail == [
        ("ISO 9001:2015", "Clause 7.5"),
        ("ISO 27001:2022", "Clause 7.5"),
        ("NDPA 2023", "Section 24"),
    ]


def test_standards_list_can_grow_beyond_three():
    """Confirmed by the officer: the standards count can grow further."""
    ctx = _sample_context(standards=[
        StandardEntry(name="ISO 9001:2015", reference="Clause 7.5"),
        StandardEntry(name="ISO 27001:2022", reference="Clause 7.5"),
        StandardEntry(name="NDPA 2023", reference="Section 24"),
        StandardEntry(name="ISO 22301:2019", reference="Clause 8.4"),
    ])
    result = merge(_build_master(), ctx)
    d = Document(io.BytesIO(result.document))
    assert len(d.tables[0].rows) == 17  # 13 fixed + 4 standards
    assert scan_markers(result.document) == []


def test_empty_standards_produces_no_extra_rows():
    result = merge(_build_master(), _sample_context(standards=[]))
    d = Document(io.BytesIO(result.document))
    assert len(d.tables[0].rows) == 13  # just the fixed rows
    assert scan_markers(result.document) == []


# -----------------------------------------------------------------------------
#  Fail-loud guarantees
# -----------------------------------------------------------------------------

def test_unknown_marker_in_template_raises_missing_context():
    # Template references {{ bogus_field }} which MergeContext never provides.
    bad = _build_master(extra_body="Stray: {{ bogus_field }}")
    with pytest.raises(MissingContextError):
        merge(bad, _sample_context())


def test_leftover_literal_marker_is_detected_and_blocks_publish():
    # {% raw %}…{% endraw %} makes jinja emit a LITERAL {{ orphan }} into the
    # output, simulating a marker docxtpl could not resolve. Must be caught.
    doc = Document()
    doc.add_paragraph("{% raw %}{{ orphan }}{% endraw %}")
    buf = io.BytesIO()
    doc.save(buf)
    with pytest.raises(UnresolvedMarkerError) as exc:
        merge(buf.getvalue(), _sample_context())
    assert any("orphan" in m for m in exc.value.markers)


def test_scan_markers_finds_tokens_in_unrendered_template():
    # The scanner itself must see markers before any render.
    markers = scan_markers(_build_master())
    joined = " ".join(markers)
    assert "{{ doc_code }}" in joined
    assert "{%tr for r in revision_history %}" in joined
    assert "{%tr for s in standards %}" in joined


# -----------------------------------------------------------------------------
#  XML-safety and Unicode (edge cases that silently corrupt naive implementations)
# -----------------------------------------------------------------------------

def test_xml_significant_characters_are_escaped_not_corrupting():
    # A title with & < > " must render as text, and the docx must stay valid
    # (Document() re-open would raise if the XML were corrupted).
    ctx = _sample_context(title='Fast & Secure <Ops> "Live"')
    result = merge(_build_master(), ctx)
    text = _all_text(result.document)          # re-opens the docx == validity check
    assert 'Fast & Secure <Ops> "Live"' in text
    assert scan_markers(result.document) == []


def test_unicode_values_preserved():
    ctx = _sample_context(final_approval="Adékúnlé Oyègbésan", title="Procédure CDI")
    result = merge(_build_master(), ctx)
    text = _all_text(result.document)
    assert "Adékúnlé Oyègbésan" in text
    assert "Procédure CDI" in text


# -----------------------------------------------------------------------------
#  render_document vs merge (validation seam)
# -----------------------------------------------------------------------------

def test_corrupt_source_raises_template_render_error():
    with pytest.raises(TemplateRenderError):
        merge(b"this is not a docx", _sample_context())


def test_control_fact_containing_marker_token_is_rejected_up_front():
    # A value that itself contains "{{" would trip the leftover scan; we reject
    # it with a precise error instead of a confusing UnresolvedMarkerError.
    ctx = _sample_context(title="Example uses {{ handlebars }}")
    with pytest.raises(MergeError) as exc:
        merge(_build_master(), ctx)
    assert "title" in str(exc.value)


def test_history_value_containing_marker_token_is_rejected():
    ctx = _sample_context(revision_history=[
        RevisionHistoryEntry(version="00", date="May 2025",
                             purpose="Documents {% weird %} syntax", approved_by="X"),
    ])
    with pytest.raises(MergeError) as exc:
        merge(_build_master(), ctx)
    assert "revision_history[0].purpose" in str(exc.value)


def test_standards_value_containing_marker_token_is_rejected():
    """The generalised list-scan (added for v06) must cover `standards` too."""
    ctx = _sample_context(standards=[StandardEntry(name="{{ hack }}", reference="x")])
    with pytest.raises(MergeError) as exc:
        merge(_build_master(), ctx)
    assert "standards[0].name" in str(exc.value)


def test_render_document_does_not_validate_but_merge_does():
    # render_document returns bytes even with a raw leftover; merge refuses.
    doc = Document()
    doc.add_paragraph("{% raw %}{{ orphan }}{% endraw %}")
    buf = io.BytesIO()
    doc.save(buf)
    src = buf.getvalue()
    rendered = render_document(src, _sample_context())   # no raise
    assert scan_markers(rendered)                          # leftover present
    with pytest.raises(UnresolvedMarkerError):
        merge(src, _sample_context())
