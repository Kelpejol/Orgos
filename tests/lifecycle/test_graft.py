# =============================================================================
# tests/lifecycle/test_graft.py — attaching a CDT cover onto an existing doc
#
# graft_cover_onto_document() is pure (no I/O) — every test builds a small
# fake "master template" and a small fake "existing document" directly with
# python-docx, rather than depending on the real, live-uploaded master
# templates (kept isolated, house pattern — see test_merge.py).
# =============================================================================

import io

from docx import Document
from docx.oxml.ns import qn

from lifecycle.graft import graft_cover_onto_document
from lifecycle.merge import scan_markers


def _fake_master_template() -> bytes:
    """
    Minimal but structurally faithful stand-in for a real master template:
    a cover table with markers, then the "[Body content begins here]"
    placeholder paragraph, then the document's sectPr — same shape
    scripts/build_cdt_master_templates.py produces.
    """
    doc = Document()
    doc.add_paragraph("DRAGNET SOLUTIONS")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Document Code"
    table.rows[0].cells[1].text = "{{ doc_code }}"
    table.rows[1].cells[0].text = "Version"
    table.rows[1].cells[1].text = "{{ version }}"
    doc.add_page_break()
    doc.add_paragraph("[Body content begins here — 1.0 Purpose & Scope, etc.]")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _fake_existing_document(*, with_image: bool = False) -> bytes:
    doc = Document()
    doc.add_heading("Legacy Document Title", level=1)
    doc.add_paragraph("Plain body paragraph, written with no CDT cover at all.")
    p = doc.add_paragraph()
    p.add_run("Bold bit ").bold = True
    p.add_run("and italic bit.").italic = True
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "left"
    table.rows[0].cells[1].text = "right"
    if with_image:
        import os
        logo_path = os.path.join(
            os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))),
            "assets", "branding", "dragnet_logo.png",
        )
        doc.add_picture(logo_path)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_graft_preserves_cover_markers():
    result = graft_cover_onto_document(_fake_existing_document(), _fake_master_template())
    markers = scan_markers(result.document)
    assert "{{ doc_code }}" in markers
    assert "{{ version }}" in markers


def test_graft_carries_over_text_and_formatting():
    result = graft_cover_onto_document(_fake_existing_document(), _fake_master_template())
    doc = Document(io.BytesIO(result.document))
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Legacy Document Title" in all_text
    assert "Plain body paragraph, written with no CDT cover at all." in all_text

    bold_run = next(
        r for p in doc.paragraphs for r in p.runs if r.text == "Bold bit "
    )
    assert bold_run.bold is True


def test_graft_carries_over_tables_in_order():
    result = graft_cover_onto_document(_fake_existing_document(), _fake_master_template())
    doc = Document(io.BytesIO(result.document))
    # First table is the master's cover table, second is the existing doc's.
    assert len(doc.tables) == 2
    assert doc.tables[0].rows[0].cells[0].text == "Document Code"
    assert doc.tables[1].rows[0].cells[0].text == "left"


def test_graft_drops_placeholder_paragraph():
    result = graft_cover_onto_document(_fake_existing_document(), _fake_master_template())
    doc = Document(io.BytesIO(result.document))
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Body content begins here" not in all_text


def test_graft_carries_over_image():
    result = graft_cover_onto_document(
        _fake_existing_document(with_image=True), _fake_master_template()
    )
    doc = Document(io.BytesIO(result.document))
    assert result.images_carried == 1
    assert result.images_skipped == 0
    # The image's relationship id must resolve to a real, readable part.
    found_blip = False
    for p in doc.paragraphs:
        for run in p.runs:
            for blip in run._r.iter(qn("a:blip")):
                rid = blip.get(qn("r:embed"))
                part = doc.part.related_parts[rid]
                assert len(part.blob) > 0
                found_blip = True
    assert found_blip


def test_graft_on_document_with_no_tables_or_images():
    doc = Document()
    doc.add_paragraph("Just one plain paragraph, nothing else.")
    buf = io.BytesIO()
    doc.save(buf)

    result = graft_cover_onto_document(buf.getvalue(), _fake_master_template())
    assert result.images_carried == 0
    assert result.images_skipped == 0
    out_doc = Document(io.BytesIO(result.document))
    all_text = "\n".join(p.text for p in out_doc.paragraphs)
    assert "Just one plain paragraph, nothing else." in all_text
