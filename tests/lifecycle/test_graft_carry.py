# =============================================================================
# tests/lifecycle/test_graft_carry.py — what a graft must carry, not degrade
#
# Measured on the real library: 85 of 123 documents use styles the template
# lacks (Table, Compact, FirstParagraph) and 52 use automatic numbering. A
# graft that drops either silently changes how a document reads.
# =============================================================================

import io

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from lifecycle.graft import graft_cover_onto_document, unresolved_relationships
from tests.lifecycle.test_graft import _fake_master_template


def _bytes(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _numbered_original() -> bytes:
    doc = Document()
    style = doc.styles.add_style("Compact", WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles["Normal"]
    numbering = doc.part.numbering_part.element
    numbering.insert(0, parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="41"><w:nsid w:val="ABCD1234"/>'
        '<w:multiLevelType w:val="singleLevel"/>'
        '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl>'
        '</w:abstractNum>'))
    numbering.append(parse_xml(
        f'<w:num {nsdecls("w")} w:numId="77"><w:abstractNumId w:val="41"/></w:num>'))
    for text in ("First step", "Second step"):
        p = doc.add_paragraph(text, style="Compact")
        p._p.get_or_add_pPr().append(parse_xml(
            f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="77"/></w:numPr>'))
    return _bytes(doc)


def test_a_style_the_template_lacks_is_carried_not_degraded():
    result = graft_cover_onto_document(_numbered_original(), _fake_master_template())
    out = Document(io.BytesIO(result.document))
    assert "Compact" in {s.style_id for s in out.styles}
    assert result.styles_carried >= 1


def test_numbered_steps_keep_their_numbering():
    result = graft_cover_onto_document(_numbered_original(), _fake_master_template())
    out = Document(io.BytesIO(result.document))
    steps = [p for p in out.paragraphs if p.text in ("First step", "Second step")]
    ids = {p._p.pPr.numPr.numId.val for p in steps}
    assert len(ids) == 1                                    # still one list
    numbering = out.part.numbering_part.element
    nums = {n.get(qn("w:numId")): n for n in numbering.iter(qn("w:num"))}
    (new_id,) = {str(i) for i in ids}
    assert new_id in nums                                   # the id points at a real definition…
    abstract_id = nums[new_id].find(qn("w:abstractNumId")).get(qn("w:val"))
    abstracts = {a.get(qn("w:abstractNumId")) for a in numbering.iter(qn("w:abstractNum"))}
    assert abstract_id in abstracts                         # …which exists…
    fmt = [a for a in numbering.iter(qn("w:abstractNum")) if a.get(qn("w:abstractNumId")) == abstract_id][0]
    assert fmt.find(".//" + qn("w:numFmt")).get(qn("w:val")) == "decimal"   # …and still counts 1. 2. 3.
    assert result.numbering_carried == 1


def test_abstract_numbering_comes_before_numbers_in_the_numbering_part():
    # Word rejects a numbering part whose <w:num> precede an <w:abstractNum>.
    result = graft_cover_onto_document(_numbered_original(), _fake_master_template())
    numbering = Document(io.BytesIO(result.document)).part.numbering_part.element
    tags = [c.tag for c in numbering]
    last_abstract = max(i for i, t in enumerate(tags) if t == qn("w:abstractNum"))
    first_num = min(i for i, t in enumerate(tags) if t == qn("w:num"))
    assert last_abstract < first_num


def test_a_copied_numbering_definition_does_not_reuse_the_originals_identity():
    result = graft_cover_onto_document(_numbered_original(), _fake_master_template())
    numbering = Document(io.BytesIO(result.document)).part.numbering_part.element
    assert not [n for n in numbering.iter(qn("w:nsid")) if n.get(qn("w:val")) == "ABCD1234"]


def test_a_hyperlink_the_graft_cannot_carry_is_reported_not_left_dangling():
    doc = Document()
    p = doc.add_paragraph("See ")
    p._p.append(parse_xml(
        f'<w:hyperlink {nsdecls("w")} xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'r:id="rId999"><w:r><w:t>this link</w:t></w:r></w:hyperlink>'))
    result = graft_cover_onto_document(_bytes(doc), _fake_master_template())
    assert any("rId999" in u for u in result.unresolved)


def test_a_clean_graft_reports_nothing_unresolved():
    result = graft_cover_onto_document(_numbered_original(), _fake_master_template())
    assert result.unresolved == ()


def test_skip_leading_leaves_out_the_originals_header_block():
    doc = Document()
    doc.add_paragraph("Old title line")
    doc.add_table(rows=1, cols=2).rows[0].cells[0].text = "Document code"
    doc.add_paragraph("Body that must stay")
    result = graft_cover_onto_document(_bytes(doc), _fake_master_template(), skip_leading=2)
    texts = [p.text for p in Document(io.BytesIO(result.document)).paragraphs]
    assert "Old title line" not in texts and "Body that must stay" in texts
    assert result.skipped_leading == 2 and result.cover_elements > 0
