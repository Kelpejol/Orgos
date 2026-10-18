# =============================================================================
# tests/lifecycle/test_cover_detect.py
#
# "Attach CDT cover" must only be offered to a document that has no cover.
# Getting this wrong in either direction is costly: a false "no cover" grafts a
# second cover onto a document that already has one, and a false "has cover"
# leaves a document stranded off the template.
# =============================================================================

import io

from docx import Document

from lifecycle.cover_detect import detect_cover


def _build(cover_rows=None, history=False, body=("Body text.",), markers=False):
    d = Document()
    if cover_rows:
        t = d.add_table(rows=len(cover_rows), cols=2)
        for i, (label, value) in enumerate(cover_rows):
            t.cell(i, 0).text = label
            t.cell(i, 1).text = value
    if history:
        h = d.add_table(rows=2, cols=4)
        for i, head in enumerate(["Version", "Date", "Purpose of Change", "Approved By"]):
            h.cell(0, i).text = head
    for p in body:
        d.add_paragraph("{{ doc_code }}" if markers else p)
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


_FULL_COVER = [
    ("Document Title", "Access Policy"),
    ("Document Code", "DRG-ISMS-POL-ACC"),
    ("Classification", "Internal Use Only"),
    ("Owner", "Head of IT"),
    ("Version", "1.0"),
]


def test_live_markers_count_as_a_cover():
    r = detect_cover(_build(markers=True))
    assert r.has_cover and r.source == "markers"
    assert r.as_dict()["marker_count"] >= 1


def test_merged_cover_table_is_recognised_without_markers():
    r = detect_cover(_build(_FULL_COVER, history=True))
    assert r.has_cover and r.source == "cover table"
    assert "document code" in r.matched_labels
    assert r.has_revision_history


def test_plain_document_has_no_cover():
    r = detect_cover(_build())
    assert not r.has_cover and r.source == "none"
    assert "attaching one" in r.detail.lower()


def test_an_ordinary_two_column_table_is_not_a_cover():
    r = detect_cover(_build([("Owner", "Head of IT"), ("Frequency", "Monthly")]))
    assert not r.has_cover


def test_older_hand_built_labels_still_count():
    r = detect_cover(_build([
        ("Title", "Access Policy"),
        ("Document Reference", "DRG-ISMS-POL-ACC"),
        ("Document Owner", "Head of IT"),
        ("Version No", "1.0"),
    ]))
    assert r.has_cover and r.source == "cover table"


def test_a_pdf_is_reported_as_unreadable_not_as_missing():
    r = detect_cover(b"%PDF-1.7 not a docx", "Policy.pdf")
    assert r.source == "unreadable" and not r.has_cover
    assert "PDF" in r.detail


def test_unreadable_bytes_do_not_raise():
    r = detect_cover(b"junk", "Policy.docx")
    assert r.source == "unreadable" and not r.has_cover
