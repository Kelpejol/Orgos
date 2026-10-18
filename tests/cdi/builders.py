# =============================================================================
# tests/cdi/builders.py — build small Word documents in the CDI shapes
#
# Each test states the exact structure it is about, rather than leaning on a
# real document that could change underneath it. The real documents get their
# own golden tests in test_real_documents.py.
# =============================================================================

from __future__ import annotations

import io
from typing import Iterable, Optional

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

Row = Iterable[str]


class DocBuilder:
    def __init__(self) -> None:
        self.doc = Document()
        self._change_id = 0

    # ── content ──────────────────────────────────────────────────────────────

    def para(self, text: str, style: Optional[str] = None) -> "DocBuilder":
        if style and style not in [s.name for s in self.doc.styles]:
            # The blank template lacks some built-ins Word documents use,
            # e.g. "toc 1" for table-of-contents entries.
            self.doc.styles.add_style(style, WD_STYLE_TYPE.PARAGRAPH)
        self.doc.add_paragraph(text, style=style)
        return self

    def heading(self, text: str, level: int = 1) -> "DocBuilder":
        return self.para(text, style=f"Heading {level}")

    def table(self, rows: list[Row]) -> "DocBuilder":
        rows = [list(r) for r in rows]
        width = max(len(r) for r in rows)
        t = self.doc.add_table(rows=len(rows), cols=width)
        for i, r in enumerate(rows):
            for j, value in enumerate(r):
                t.cell(i, j).text = value
        return self

    def control_block(self, **fields: str) -> "DocBuilder":
        """A CDI control block table: a title row, then Label | Value rows."""
        labels = {
            "title": "Document Title", "code": "Document Code", "layer": "Layer",
            "domain": "Domain", "parent": "Parent Document", "children": "Child Documents",
            "owner": "Document Owner", "approval": "Approval Authority",
            "classification": "Classification", "status": "Issue Status",
            "standards": "Standards", "implemented": "Standard controls implemented",
            "version": "Version",
        }
        rows = [["Control block", ""]] + [[labels[k], v] for k, v in fields.items()]
        return self.table(rows)

    def control(self, control_id: str, title: str, **fields: str) -> "DocBuilder":
        """A control heading followed by its labelled lines."""
        labels = {
            "requirement": "Requirement", "trigger": "Trigger",
            "responsibility": "Accountable / Executing", "evidence": "Evidence",
            "risk": "Risk", "failure_route": "Failure route",
            "exception_route": "Exception route", "mapping": "Mapping",
        }
        self.heading(f"{control_id} {title}", level=3)
        for key, value in fields.items():
            self.para(f"{labels.get(key, key)}: {value}")
        return self

    # ── tracked changes ─────────────────────────────────────────────────────

    def _tracked(self, paragraph, text: str, kind: str) -> None:
        self._change_id += 1
        wrap = OxmlElement(f"w:{kind}")
        wrap.set(qn("w:id"), str(self._change_id))
        wrap.set(qn("w:author"), "Reviewer")
        wrap.set(qn("w:date"), "2026-10-01T00:00:00Z")
        run = OxmlElement("w:r")
        t = OxmlElement("w:delText" if kind == "del" else "w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = text
        run.append(t)
        wrap.append(run)
        paragraph._p.append(wrap)

    def para_with_changes(self, kept: str, inserted: str = "", deleted: str = "") -> "DocBuilder":
        p = self.doc.add_paragraph(kept)
        if deleted:
            self._tracked(p, deleted, "del")
        if inserted:
            self._tracked(p, inserted, "ins")
        return self

    def table_row_inserted(self, label: str, value: str) -> "DocBuilder":
        """Append a two-cell table whose every cell's text is a tracked insertion."""
        t = self.doc.add_table(rows=1, cols=2)
        for cell, text in zip(t.rows[0].cells, (label, value)):
            self._tracked(cell.paragraphs[0], text, "ins")
        return self

    # ── output ───────────────────────────────────────────────────────────────

    def bytes(self) -> bytes:
        out = io.BytesIO()
        self.doc.save(out)
        return out.getvalue()


def defining_standard(**overrides) -> DocBuilder:
    """A small Standard in the CDI v06 shape, with two complete controls."""
    b = DocBuilder()
    b.control_block(
        title="Test Standard", code="DRG-QI-STD-TST", layer="Layer 3 Standard",
        domain="Quality & Compliance", parent="Test Policy (DRG-QI-POL-TST)",
        owner="Compliance Executive (QMS & NDPA)",
        approval="First Pass: Chief Compliance Officer. Final Approval: Managing Director / CEO (route)",
        classification="Internal Use Only. Distribution: All Staff",
        standards="ISO/IEC 27001:2022 Clause 7.5, Annex A 5.1",
    )
    b.heading("1.0 Purpose").para("1.1 This Standard sets the test controls.")
    b.heading("5.0 Controls")
    b.control(
        "TST-001", "Approved parent",
        requirement="Every document MUST link to an approved parent.",
        trigger="Creation, submission.",
        responsibility="Compliance Executive (QMS & NDPA) / platform check V03.",
        evidence="REC-0623 Controlled Document Register; validation result REC-0635.",
        risk="A document is issued without authority.",
        failure_route="Submission blocked.",
        exception_route="None.",
        mapping="ISO 9001 7.5.2; ISO/IEC 27001 7.5.2, A.5.1.",
    )
    b.control(
        "TST-002", "Identity and version",
        requirement="Every document MUST carry a unique code.",
        trigger="Issue, revision.",
        responsibility="Document Owner.",
        evidence="REC-0624 Document Change Request.",
        risk="Staff use the wrong version.",
        failure_route="Submission blocked.",
        exception_route="None.",
        mapping="ISO 9001 7.5.3.",
    )
    b.heading("6.0 Roles")
    b.table([["Role", "Responsibility under this Standard"],
             ["Managing Director / CEO", "Final Approval"],
             ["Document Owner", "Owns the document"]])
    return b


def implementing_procedure() -> DocBuilder:
    """A small Procedure that carries out two controls, BCDR style."""
    b = DocBuilder()
    b.control_block(
        title="Test Procedure", code="DRG-ISMS-PRO-TST", layer="Layer 4 Procedure",
        parent="Test Standard (DRG-ISMS-STD-TST)",
        children="Continuity Register (REC-0675); Business Impact Assessment (REC-0676)",
        owner="Head of IT and Infrastructure",
        implemented="TST-01; TST-02",
    )
    b.heading("3.0 Roles")
    b.table([["Role", "Duty"],
             ["Accountable Department Head", "Owns a process."],
             ["Head of IT and Infrastructure", "Keeps the Continuity Register."]])
    b.heading("4.0 Business impact assessment (TST-01)")
    b.para("4.1 The Accountable Department Head records each process in REC-0676.")
    b.para("4.2 The assessment is reviewed each year.")
    b.para("Risk statement (TST-01): A process could stay down too long.")
    b.heading("5.0 Continuity Register (implements TST-02)")
    b.para("5.1 The Head of IT and Infrastructure keeps one Continuity Register (REC-0675).")
    b.heading("Annex A. Control map")
    b.table([["Standard control", "Clause in this Procedure"],
             ["TST-01 Business impact assessment", "4.0"],
             ["TST-02 Continuity register", "5.0"]])
    b.heading("Annex B. Records")
    b.table([["Record", "Created by", "At step", "Kept in", "Retention"],
             ["REC-0675 Continuity Register", "Head of IT and Infrastructure", "5.1", "Catalogue folder", "Annex C"],
             ["REC-0676 Business impact assessment", "Accountable Department Head", "4.1", "Catalogue folder", "Annex C"]])
    b.heading("Annex C. Open parameters")
    b.table([["Parameter", "Used in", "Value"],
             ["Maximum tolerable downtime", "4.1", ""],
             ["Review month", "4.2", "January"]])
    return b
