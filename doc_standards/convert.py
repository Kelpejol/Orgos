# =============================================================================
# doc_standards/convert.py — convert a document onto the CDT, provably
#
# A v06 document opens with a hand-typed header block: a title line and a table
# of facts (code, title, owner, parent, approvers, version, dates…). Those
# facts are typed into the file, so the file and the Document Register can
# disagree — the drift. Converting replaces the block with the CDT cover, whose
# values are filled from the Register, and keeps every fact the block stated.
#
#   1. READ    the header block with the CDI parser (what the document says
#              about itself — never a guess from the file name or folder).
#   2. SAVE    every label/value verbatim in the cover facts (`extras`).
#   3. REMOVE  only the rows the cover regenerates (title, code, layer, domain,
#              parent, owner, approvers, classification, distribution, version,
#              dates, status). Every OTHER row — controls implemented,
#              supersedes, related and child documents, standards, review
#              cycle — stays in the document, verbatim, right after the cover:
#              the cover has no cell for them, and OrgOS reads them from the
#              file. Dropping them would silently change what the document says
#              it implements.
#   4. REBUILD the file: template cover, then the body UNCHANGED.
#   5. PROVE   it. Nothing is offered unless every proof holds:
#                body       the body text is identical, element by element
#                reads-same OrgOS's own reader finds the same code, title,
#                           relationships and control definitions after as
#                           before (the version, dates and status excepted)
#                removed    the only rows taken out are ones the cover regenerates
#                sound      no reference to a part of the original that no
#                           longer exists; numbering definitions all resolve
#                merges     the cover fills with no marker left over
#
# What it will not do is invent: a fact the document does not state stays
# blank, and a document that does not open with a recognisable header block is
# refused with the reason, not guessed at.
#
# Pure: bytes in, bytes and findings out. No I/O. The router fetches the
# master template and does all SharePoint work.
# =============================================================================

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Optional

from docx import Document
from docx.oxml.ns import qn

from cdi.parser import header_key, parse_document
from lifecycle.graft import graft_cover_onto_document, unresolved_relationships
from lifecycle.merge import merge, scan_markers
from lifecycle.schemas import CdtCoverFacts, MergeContext, StandardEntry

_SECT_PR = qn("w:sectPr")
_P, _TBL, _TR, _TC = qn("w:p"), qn("w:tbl"), qn("w:tr"), qn("w:tc")
_TEXT_TAGS = (qn("w:t"), qn("w:delText"))
# Invisible markers Word puts between body elements. They never count as content.
_MARKERS = {qn("w:bookmarkStart"), qn("w:bookmarkEnd"), qn("w:proofErr"),
            qn("w:permStart"), qn("w:permEnd")}
_CLOSERS = {qn("w:bookmarkEnd"), qn("w:proofErr"), qn("w:permEnd")}
_BM_ID, _BM_NAME, _ANCHOR = qn("w:id"), qn("w:name"), qn("w:anchor")

# Document-type code → the master template that carries it. Types with no
# template (registers, forms, checklists, playbooks) are refused: a register is
# a record, not a controlled document with a cover.
TEMPLATE_BY_TYPE = {
    "STD": "Standard", "POL": "Policy", "PRO": "Procedure", "SPEC": "Procedure",
    "GDL": "Guideline", "MAN": "Manual",
}

# Header-block rows the cover regenerates, and so are removed from the file:
#   covered   the template cover has a cell for them
#   dropped   the Register holds them (status), or approval stamps them (version, dates)
_COVERED = {"title", "code", "layer", "domain", "parent", "owner", "first_pass", "final_approval",
            "classification", "distribution", "version", "effective_date", "next_review"}
_STATUS_LABELS = {"issue status", "status", "document status"}
_REGENERATED = {"version": "Version", "effective_date": "Effective date", "next_review": "Next review due"}
# What a reader of the converted file must still find, exactly as before.
_MUST_READ_SAME = ("code", "title", "type_code", "layer_as_written", "domain", "parent_as_written",
                   "owner", "first_pass_approval", "final_approval", "classification", "distribution",
                   "review_cycle", "supersedes", "legacy_code", "approval_authority",
                   "child_documents_as_written", "related_as_written", "standards_as_written",
                   "controls_implemented", "controls_implemented_as_written")


class ConversionError(Exception):
    """The document cannot be converted; the message says why, for a person."""


# ── reading the original ────────────────────────────────────────────────────

def _text(el) -> str:
    return "".join(t.text or "" for t in el.iter(*_TEXT_TAGS))


def _norm(text: str) -> str:
    return re.sub(r"\s+", " ", (text or "").strip().lower())


def _signature(elements) -> list:
    """What a reader sees of each body element — paragraph text; for a table,
    its rows and cells. Used on both sides of the conversion, so equality means
    nothing the reader sees has changed."""
    out = []
    for el in elements:
        if el.tag in _MARKERS:
            continue                                   # invisible: not something a reader sees
        if el.tag == _TBL:
            rows = ["\t".join(_text(tc) for tc in tr.findall(_TC)) for tr in el.iter(_TR)]
            out.append(("tbl", "\n".join(rows)))
        else:
            out.append((el.tag.split("}")[-1], _text(el)))
    return out


def _body_children(doc) -> list:
    return [c for c in doc.element.body if c.tag != _SECT_PR]


def _table_rows(tbl) -> list[list[str]]:
    rows = []
    for tr in tbl.iter(_TR):
        cells = [_text(tc).strip() for tc in tr.findall(_TC)]
        rows.append(list(dict.fromkeys(cells)))      # a merged cell repeats itself
    return rows


def _is_header_table(rows: list[list[str]]) -> list[tuple[str, str]]:
    """The label/value pairs, when this table opens like a control block; else []."""
    if not rows or not rows[0]:
        return []
    first = _norm(" ".join(rows[0]))
    opens = first in ("control block", "control metadata value", "field value", "field detail") \
        or bool(header_key(rows[0][0]))
    if not opens:
        return []
    pairs = [(r[0], " ".join(r[1:]).strip()) for r in rows if r and header_key(r[0])]
    return pairs if len(pairs) >= 4 else []


def _front_matter_ok(style: str, text: str, title: str) -> bool:
    """May this paragraph, sitting above the header table, be removed with it?
    Only a title or a status/brand line — never a sentence of the document."""
    t = _norm(text)
    if len(text) > 200:
        return False
    if style.lower() in ("title", "subtitle"):
        return True
    nt = _norm(title)
    # The title line itself, or the title with a short decoration ("Policy — Draft").
    # Never a sentence that merely contains the title's letters: that is document content.
    if nt and (t == nt or (len(nt) >= 8 and nt in t and len(t) <= len(nt) + 60)):
        return True
    return t.startswith(("dragnet", "draft", "working draft", "confidential", "document control"))


@dataclass
class HeaderBlock:
    front: int                                  # body elements above the first header table (title, markers)
    removed_text: list[str]                     # the title/status paragraphs among them
    table_indices: list[int]                    # body positions of the header table(s)
    tables: list[list[tuple[str, str]]]         # label/value pairs, per table
    orphan_ends: frozenset = frozenset()        # bookmark ids whose START is removed with the front matter


def find_header_block(doc, title: str) -> HeaderBlock:
    children = _body_children(doc)
    removed: list[str] = []
    indices: list[int] = []
    tables: list[list[tuple[str, str]]] = []
    pending: list[str] = []
    front = None
    for i, el in enumerate(children):
        if el.tag == _P:
            text = _text(el).strip()
            if not text:
                continue
            if tables:                                # past the block: first real paragraph
                break
            style_id = ""
            ppr = el.find(qn("w:pPr"))
            if ppr is not None and ppr.find(qn("w:pStyle")) is not None:
                style_id = ppr.find(qn("w:pStyle")).get(qn("w:val")) or ""
            try:
                style_name = next((s.name for s in doc.styles if s.style_id == style_id), style_id)
            except Exception:
                style_name = style_id
            if len(pending) >= 3 or not _front_matter_ok(style_name, text, title):
                raise ConversionError(
                    f"The document opens with text that is not a title or status line "
                    f"(“{text[:70]}…”), so OrgOS cannot tell where its header block is.")
            pending.append(text)
        elif el.tag in _MARKERS:
            continue                                  # invisible; sits wherever Word put it
        elif el.tag == _TBL:
            pairs = _is_header_table(_table_rows(el))
            if not pairs:
                if tables:
                    break
                raise ConversionError(
                    "The first table in the document is not a header block (a table of "
                    "Document code, Owner, Parent… facts).")
            if front is None:
                front = i
                removed, pending = pending, []
            indices.append(i)
            tables.append(pairs)
        else:
            break
    if not tables:
        raise ConversionError(
            "No header block found at the top of the document. Its facts may be written as "
            "'Label: value' lines, which cannot be converted automatically yet.")
    orphans = _bookmarks_left_open(children, front)
    return HeaderBlock(front=front, removed_text=removed, table_indices=indices, tables=tables,
                       orphan_ends=orphans)


def _bookmarks_left_open(children: list, skip: int) -> frozenset:
    """
    Removing the front matter takes its bookmarks with it. A section bookmark
    opens at the title and closes pages later, so its END is then left behind
    without a start — returned here so it can be removed too. Safe only when
    nothing in the body points at a bookmark being removed; if something does,
    refuse.
    """
    gone, rest = children[:skip], children[skip:]
    inside = [node for el in gone for node in el.iter()]       # a bookmark may open or close within a paragraph
    opened = {node.get(_BM_ID): node.get(_BM_NAME) for node in inside if node.tag == qn("w:bookmarkStart")}
    closed = {node.get(_BM_ID) for node in inside if node.tag == qn("w:bookmarkEnd")}
    if closed - set(opened):
        raise ConversionError("A bookmark in the header area opens before it, so the header block "
                              "cannot be removed without damaging it.")
    orphans = frozenset(set(opened) - closed)
    names = {n for n in opened.values() if n}
    if not names:
        return orphans
    for el in rest:
        for node in el.iter():
            if node.get(_ANCHOR) in names or any(
                    n in (node.text or "") for n in names if node.tag == qn("w:instrText")):
                raise ConversionError(
                    "Something in the body (a link, a cross-reference or a table-of-contents entry) points "
                    "at the document's title area, which conversion replaces. It would be left broken.")
    return orphans


def _row_is_regenerated(label: str) -> bool:
    key = header_key(label)
    return key in _COVERED or _norm(label) in _STATUS_LABELS


def _strip_regenerated_rows(table_el) -> tuple[list, bool]:
    """
    Remove from a header table the rows the cover regenerates. Returns (labels
    removed, whether the table now has nothing left worth keeping). A caption
    row ("Control block") is kept only if some other row is.
    """
    removed, kept_facts = [], 0
    captions = []
    for tr in list(table_el.iter(_TR)):
        cells = list(dict.fromkeys(_text(tc).strip() for tc in tr.findall(_TC)))
        label = cells[0] if cells else ""
        if not label and len(cells) <= 1:
            captions.append(tr)                       # an empty spacer row
        elif len(cells) <= 2 and not header_key(label) and _norm(label) in (
                "control block", "control metadata value", "field value", "field detail") \
                and not (len(cells) == 2 and cells[1]):
            captions.append(tr)
        elif _row_is_regenerated(label):
            removed.append((label, " ".join(cells[1:]).strip()))
            tr.getparent().remove(tr)
        else:
            kept_facts += 1
    if kept_facts == 0:
        return removed, True
    return removed, False


def _page_furniture(doc) -> list[str]:
    """Text in the original's page headers and footers — which the template replaces."""
    seen: list[str] = []
    for section in doc.sections:
        for part in (section.header, section.first_page_header, section.footer, section.first_page_footer):
            for p in part.paragraphs:
                if p.text.strip() and p.text.strip() not in seen:
                    seen.append(p.text.strip())
    return seen


# ── the plan ────────────────────────────────────────────────────────────────

@dataclass
class Proof:
    name: str
    ok: bool
    detail: str

    def as_dict(self) -> dict:
        return {"name": self.name, "ok": self.ok, "detail": self.detail}


@dataclass
class ConversionPlan:
    ok: bool = False
    reason: str = ""                                   # why not, when not ok
    filename: str = ""
    code: str = ""
    title: str = ""
    doc_type: str = ""                                 # the template used
    removed: list[str] = field(default_factory=list)   # text removed from the top
    block_rows: list[tuple[str, str]] = field(default_factory=list)
    taken_out: list[tuple[str, str]] = field(default_factory=list)    # rows the cover regenerates
    stays: list[tuple[str, str]] = field(default_factory=list)        # rows kept, in place, verbatim
    cover: list[tuple[str, str]] = field(default_factory=list)       # what the cover will say
    kept_in_register: list[tuple[str, str]] = field(default_factory=list)
    regenerated: list[tuple[str, str, str]] = field(default_factory=list)  # label, typed value, why
    furniture_replaced: list[str] = field(default_factory=list)
    carried: dict = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)
    proofs: list[Proof] = field(default_factory=list)
    facts: Optional[CdtCoverFacts] = None
    source_bytes: bytes = b""          # the converted file, cover markers still live
    preview_bytes: bytes = b""         # the same, cover filled — what a reader will see

    def as_dict(self) -> dict:
        return {
            "ok": self.ok, "reason": self.reason, "file": self.filename,
            "code": self.code, "title": self.title, "template": self.doc_type,
            "removed_from_top": self.removed,
            "block_rows": [{"label": k, "value": v} for k, v in self.block_rows],
            "taken_out": [{"label": k, "value": v} for k, v in self.taken_out],
            "stays_in_document": [{"label": k, "value": v} for k, v in self.stays],
            "cover": [{"label": k, "value": v} for k, v in self.cover],
            "kept_in_register": [{"label": k, "value": v} for k, v in self.kept_in_register],
            "regenerated": [{"label": k, "typed": v, "why": w} for k, v, w in self.regenerated],
            "page_header_footer_replaced": self.furniture_replaced,
            "carried": self.carried, "warnings": self.warnings,
            "proofs": [p.as_dict() for p in self.proofs],
        }


def template_for(type_code: str) -> str:
    name = TEMPLATE_BY_TYPE.get((type_code or "").upper())
    if not name:
        raise ConversionError(
            f"There is no master template for document type “{type_code or 'unknown'}”. "
            "A register, form or checklist is a record, not a controlled document with a cover.")
    return name


def _type_code(parsed) -> str:
    code = parsed.header.type_code
    if code:
        return code
    raise ConversionError(
        "The document states no document code, so its type is unknown. Set the code in its "
        "header block first; OrgOS will not guess a type from the file name.")


def _cover_facts(parsed, existing: Optional[CdtCoverFacts]) -> tuple[CdtCoverFacts, list]:
    h = parsed.header
    regenerated = []
    for key, label in _REGENERATED.items():
        typed = getattr(h, key, "")
        if typed:
            regenerated.append((label, typed, "set by OrgOS when the document is approved"))
    if h.issue_status:
        regenerated.append(("Issue status", h.issue_status, "the Register holds the status"))
    keep = existing or CdtCoverFacts()
    facts = CdtCoverFacts(
        domain=h.domain, parent_document=h.parent_as_written, type_layer=h.layer_as_written,
        classification=h.classification, distribution=h.distribution, owner=h.owner,
        first_pass_approval=h.first_pass_approval, final_approval=h.final_approval,
        # A draft has no version or dates. A revision already carries the ones the
        # Register holds — those are kept, never overwritten by what was typed.
        version=keep.version or "Draft",
        effective_date=keep.effective_date, next_review_due=keep.next_review_due,
        standards=[StandardEntry(name=r.standard, reference=r.clause) for r in h.standards],
        extras=dict(keep.extras),
    )
    if h.standards_as_written:
        facts.extras.setdefault("Standards (as written)", h.standards_as_written)
    return facts, regenerated


def _controls_view(parsed) -> list:
    return [(c.control_id, c.title, c.requirement, c.trigger,
             c.responsibility.as_written if c.responsibility else "",
             c.evidence_as_written, c.risk, c.failure_route, c.exception_route,
             c.mapping_as_written, c.authority) for c in parsed.controls]


def convert_document(file_bytes: bytes, filename: str, master_bytes: Optional[bytes] = None,
                     existing_cover: Optional[CdtCoverFacts] = None) -> ConversionPlan:
    """
    Plan — and, when `master_bytes` is given, build and prove — the conversion.
    Never raises for a document that cannot be converted: the plan says why.
    """
    plan = ConversionPlan(filename=filename)
    try:
        return _convert(plan, file_bytes, filename, master_bytes, existing_cover)
    except ConversionError as exc:
        plan.ok, plan.reason = False, str(exc)
        return plan
    except Exception as exc:        # a document we could not even read is refused, not guessed at
        plan.ok = False
        plan.reason = f"The document could not be read ({type(exc).__name__}: {exc})."
        return plan


def _convert(plan, file_bytes, filename, master_bytes, existing_cover) -> ConversionPlan:
    if not filename.lower().endswith(".docx"):
        raise ConversionError("Only Word (.docx) documents can be converted. Save a PDF or .doc as .docx first.")
    if scan_markers(file_bytes):
        raise ConversionError("This document is already on the CDT: it still carries {{ }} placeholders.")

    parsed = parse_document(file_bytes, filename)
    h = parsed.header
    if not h.title:
        raise ConversionError("The document states no title, so there is nothing to put on the cover.")
    if not h.code:
        raise ConversionError("The document states no document code. Set it in its header block first; "
                              "OrgOS will not take one from the file name or from a cited document.")
    plan.code, plan.title = h.code, h.title
    plan.doc_type = template_for(_type_code(parsed))

    original = Document(io.BytesIO(file_bytes))
    block = find_header_block(original, h.title)
    plan.removed = block.removed_text
    plan.block_rows = [row for table in block.tables for row in table]
    plan.furniture_replaced = _page_furniture(original)

    facts, plan.regenerated = _cover_facts(parsed, existing_cover)
    plan.facts = facts
    plan.cover = [
        ("Document title", h.title), ("Document code", h.code),
        ("Document type / layer", facts.type_layer), ("Domain", facts.domain),
        ("Parent document", facts.parent_document), ("Classification", facts.classification),
        ("Distribution", facts.distribution), ("Owner", facts.owner),
        ("First pass approval", facts.first_pass_approval), ("Final approval", facts.final_approval),
        ("Version", facts.version), ("Effective date", facts.effective_date),
        ("Next review due", facts.next_review_due),
        ("Standards", "; ".join(f"{s.name} {s.reference}".strip() for s in facts.standards)),
    ]
    # Every fact the block stated, under the label the document used for it.
    for label, value in plan.block_rows:
        facts.extras[label] = value
    plan.kept_in_register = list(facts.extras.items())
    for label, value in plan.cover:
        if value and "{{" in value:
            raise ConversionError(f"“{label}” contains a template token ({{{{), which would break the cover.")

    # ── the working copy: take out only what the cover regenerates ──
    work = Document(io.BytesIO(file_bytes))
    work_children = _body_children(work)
    for index in reversed(block.table_indices):
        table_el = work_children[index]
        removed_rows, now_empty = _strip_regenerated_rows(table_el)
        plan.taken_out = removed_rows + plan.taken_out
        if now_empty:
            table_el.getparent().remove(table_el)
        else:
            plan.stays = [(r[0], " ".join(r[1:]).strip()) for r in _table_rows(table_el)
                          if r and r[0] and header_key(r[0])] + plan.stays
    for end in [n for n in work.element.body.iter(qn("w:bookmarkEnd")) if n.get(_BM_ID) in block.orphan_ends]:
        end.getparent().remove(end)
    buf = io.BytesIO()
    work.save(buf)
    work_bytes = buf.getvalue()

    if master_bytes is None:
        plan.ok = True
        plan.reason = "Preview only — no master template was supplied, so nothing was built or proved."
        return plan

    # ── build ──
    graft = graft_cover_onto_document(work_bytes, master_bytes, skip_leading=block.front)
    plan.source_bytes = graft.document
    plan.carried = {"styles": graft.styles_carried, "numbering": graft.numbering_carried,
                    "images": graft.images_carried}
    ctx = MergeContext(
        doc_code=h.code, title=h.title, doc_type=plan.doc_type, type_layer=facts.type_layer,
        domain=facts.domain, parent_document=facts.parent_document,
        classification=facts.classification, distribution=facts.distribution, owner=facts.owner,
        first_pass_approval=facts.first_pass_approval, final_approval=facts.final_approval,
        version=facts.version, effective_date=facts.effective_date,
        next_review_due=facts.next_review_due, standards=list(facts.standards))
    try:
        plan.preview_bytes = merge(graft.document, ctx).document
    except Exception as exc:
        plan.proofs.append(Proof("merges", False, f"The cover could not be filled: {exc}"))
        plan.ok, plan.reason = False, "The converted cover could not be filled in."
        plan.source_bytes = b""
        return plan

    # ── prove ──
    converted = Document(io.BytesIO(graft.document))
    carried = _body_children(converted)[graft.cover_elements:]

    before = _signature(_body_children(Document(io.BytesIO(work_bytes)))[block.front:])
    after = _signature(carried)
    plan.proofs.append(Proof(
        "body", before == after,
        f"All {len(before)} body elements are identical to the original." if before == after else
        f"The body differs from the original ({len(before)} elements before, {len(after)} after)."))

    # Nothing may be lost except what the cover regenerates: every row of the
    # original header block is either one of those, or still in the document.
    carried_pairs = set()
    for el in carried:
        if el.tag == _TBL:
            for row in _table_rows(el):
                if row and row[0]:
                    carried_pairs.add((row[0], " ".join(row[1:]).strip()))
    lost = [f"{label}: {value[:50]}" for label, value in plan.block_rows
            if not _row_is_regenerated(label) and (label, value) not in carried_pairs]
    plan.proofs.append(Proof(
        "removed", not lost,
        f"Only rows the cover regenerates were taken out ({len(plan.taken_out)}); "
        f"{len(plan.stays)} other row(s) stay in the document." if not lost
        else "Would lose: " + "; ".join(lost) + "."))

    reread = parse_document(plan.preview_bytes, filename)
    problems = [name for name in _MUST_READ_SAME
                if getattr(reread.header, name) != getattr(h, name)]
    if reread.header.version != facts.version:
        problems.append(f"version reads “{reread.header.version}”, expected “{facts.version}”")
    if _controls_view(reread) != _controls_view(parsed):
        problems.append("control definitions")
    if reread.implemented_ids != parsed.implemented_ids:
        problems.append("implemented controls")
    plan.proofs.append(Proof(
        "reads-same", not problems,
        f"OrgOS reads the same code, title, relationships, {len(parsed.controls)} control definition(s) and "
        f"{len(parsed.implemented_ids)} implemented control(s) after as before." if not problems
        else "After conversion OrgOS reads these differently: " + ", ".join(problems) + "."))

    dangling = list(graft.unresolved)
    numbering = converted.part.numbering_part.element
    defined = {n.get(qn("w:numId")) for n in numbering.iter(qn("w:num"))}
    used = {n.get(qn("w:val")) for el in carried for n in el.iter(qn("w:numId"))
            if n.get(qn("w:val")) not in (None, "0")}
    dangling += [f"numbering {u}" for u in sorted(used - defined)]
    if graft.images_skipped:
        dangling.append(f"{graft.images_skipped} image(s) that could not be copied")
    plan.proofs.append(Proof(
        "sound", not dangling,
        "No reference points at anything missing; every list definition resolves." if not dangling
        else "Would break: " + "; ".join(dangling) + "."))

    leftovers = scan_markers(plan.preview_bytes)
    plan.proofs.append(Proof(
        "merges", not leftovers,
        "The cover fills in completely." if not leftovers else f"Left unfilled: {', '.join(leftovers)}."))

    # ── things to tell the person, not failures ──
    if re.search(rb"<w:(ins|del)\b", file_bytes) or parsed.tracked_changes:
        plan.warnings.append("The document contains tracked changes. They are carried as they are, not accepted.")
    if any(_is_revision_table(t) for t in converted.tables[2:]):
        plan.warnings.append("The document has its own revision-history table. It stays in the body; the "
                             "template adds its own, filled in when the document is approved.")
    if graft.numbering_carried:
        plan.warnings.append(f"{graft.numbering_carried} numbered/bulleted list definition(s) were carried over. "
                             "Check the numbering looks right in Word.")
    if plan.furniture_replaced:
        plan.warnings.append("The original's page header/footer text is replaced by the template's "
                             "(see “Page header and footer text replaced”).")
    if not h.standards and not facts.extras.get("Standards"):
        plan.warnings.append("The document states no standards mapping on its cover.")

    plan.ok = all(p.ok for p in plan.proofs)
    if not plan.ok:
        plan.reason = "A proof failed, so the conversion is not offered: " + "; ".join(
            f"{p.name}: {p.detail}" for p in plan.proofs if not p.ok)
        plan.source_bytes = plan.preview_bytes = b""
    return plan


def _is_revision_table(table) -> bool:
    try:
        return [c.text.strip().lower() for c in table.rows[0].cells][:2] == ["version", "date"]
    except Exception:
        return False
