# =============================================================================
# doc_standards/role_fix.py — "use this role": change a role's name in a document
#
# When the role check finds a document naming something that is not in Org Roles
# and the author picks the role that was meant, the document itself is changed
# to name it (CDI 7.8: roles are named by controlled organisational title).
#
# This edits a controlled document, so it is deliberately narrow and proves itself:
#
#   WHERE   only a role FIELD is changed — a "Control owner:" line (or the older
#           Accountable / Executing / Responsible role line), a control table's
#           role column, a header table's Owner / First pass / Final approval row,
#           a Roles table's role column. The same words in running text ("the
#           Head of IT shall…") are left alone: they are prose, not a role field.
#   WHAT    only the role's own words, formatting of the run it sits in kept.
#   PROOF   afterwards every changed paragraph must be the original with the old
#           role replaced by the new one — nothing else — and nothing may have been
#           added or removed. If not, nothing is returned.
#
# Pure: bytes in, bytes out.
# =============================================================================

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Iterator

from docx import Document
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from cdi.fields import CONTROL_FIELD_ALIASES, control_field, norm as field_norm

# Labels whose value is a role.
_CONTROL_ROLE_KEYS = ("control_owner", "responsibility", "accountable", "executing")
_CONTROL_LABELS = sorted({a for k in _CONTROL_ROLE_KEYS for a in CONTROL_FIELD_ALIASES[k]}, key=len, reverse=True)
_HEADER_LABELS = ("document owner", "owner", "domain owner", "first pass approval", "first pass", "final approval")
_ROLE_TABLE_HEADERS = ("role", "roles")


class RoleFixError(Exception):
    """The change could not be made safely; the message says why."""


@dataclass
class RoleFix:
    data: bytes
    count: int = 0
    places: list[str] = field(default_factory=list)


def _flex(text: str) -> str:
    """The text as a pattern that tolerates any spacing between its words."""
    return r"\s+".join(re.escape(w) for w in text.split())


def _paragraphs(container) -> Iterator[Paragraph]:
    for p in container.paragraphs:
        yield p
    for table in container.tables:
        for row in table.rows:
            seen = set()
            for cell in row.cells:
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                yield from _paragraphs(cell)


def _tables(container) -> Iterator[Table]:
    for table in container.tables:
        yield table
        for row in table.rows:
            seen = set()
            for cell in row.cells:
                if id(cell._tc) not in seen:
                    seen.add(id(cell._tc))
                    yield from _tables(cell)


def _replace_span(paragraph: Paragraph, start: int, end: int, new: str) -> bool:
    """Replace characters [start, end) of the paragraph's text, keeping the formatting of the run
    the span begins in. False if the span does not lie in ordinary runs."""
    runs = paragraph.runs
    if "".join(r.text for r in runs) != paragraph.text:      # text held outside plain runs (tracked changes…)
        return False
    pos = 0
    first = True
    for run in runs:
        r_start, r_end = pos, pos + len(run.text)
        pos = r_end
        if r_end <= start or r_start >= end:
            continue
        keep_before = run.text[: max(0, start - r_start)]
        keep_after = run.text[max(0, end - r_start):] if r_end > end else ""
        run.text = keep_before + (new if first else "") + keep_after
        first = False
    return not first


def _fix_paragraph_lines(doc, pattern: re.Pattern, new: str, fix: RoleFix) -> None:
    for p in _paragraphs(doc):
        # Replacing can shift later spans; take the last match first.
        for m in reversed(list(pattern.finditer(p.text))):
            if _replace_span(p, m.start("value"), m.end("value"), new):
                fix.count += 1
                fix.places.append(f"“{m.group('label').strip()}” line")


def _cell_value_matches(text: str, role: str) -> str | None:
    """How the cell holds the role: the whole cell, or the role leading a legacy 'Role; other' value."""
    t = text.strip()
    if field_norm(t.rstrip(".")) == field_norm(role):
        return "whole"
    m = re.match(rf"^\s*{_flex(role)}\s*(?=[;.])", t)
    return "lead" if m else None


def _set_cell(cell: _Cell, role: str, new: str, how: str) -> bool:
    for p in cell.paragraphs:
        text = p.text
        if how == "whole":
            m = re.search(_flex(role), text)
        else:
            m = re.match(rf"^\s*({_flex(role)})", text)
        if m:
            span = m.span(1) if how == "lead" else m.span()
            return _replace_span(p, span[0], span[1], new)
    return False


def _fix_tables(doc, role: str, new: str, fix: RoleFix) -> None:
    for table in _tables(doc):
        rows = table.rows
        if not rows:
            continue
        head = [c.text.strip() for c in rows[0].cells]
        role_columns = [i for i, h in enumerate(head) if h and control_field(h, allow_keywords=True) in _CONTROL_ROLE_KEYS]
        roles_table = len(head) >= 2 and field_norm(head[0]) in _ROLE_TABLE_HEADERS
        for row in rows[1:] if (role_columns or roles_table) else rows:
            cells = row.cells
            targets = []
            if role_columns:
                targets = [(cells[i], "a control table's role column") for i in role_columns if i < len(cells)]
            elif roles_table:
                targets = [(cells[0], "a Roles table")]
            elif len(cells) >= 2:
                label = field_norm(cells[0].text)
                if label in _HEADER_LABELS or label in {field_norm(a) for a in _CONTROL_LABELS}:
                    targets = [(cells[1], f"the “{cells[0].text.strip()}” row")]
            for cell, where in targets:
                how = _cell_value_matches(cell.text, role)
                if how and _set_cell(cell, role, new, how):
                    fix.count += 1
                    fix.places.append(where)


def _texts(doc) -> list[str]:
    return [p.text for p in _paragraphs(doc)]


def replace_role(data: bytes, from_role: str, to_role: str) -> RoleFix:
    from_role, to_role = " ".join(from_role.split()), " ".join(to_role.split())
    if not from_role or not to_role:
        raise RoleFixError("Both the role written in the document and the role to use are needed.")
    if field_norm(from_role) == field_norm(to_role) and from_role == to_role:
        raise RoleFixError("That is already the role the document names.")

    doc = Document(io.BytesIO(data))
    before = _texts(doc)
    structure = (len(doc.paragraphs), len(doc.tables))
    fix = RoleFix(data=data)

    labels = "|".join(_flex(a) for a in _CONTROL_LABELS + list(_HEADER_LABELS))
    pattern = re.compile(rf"(?P<label>(?i:{labels}))\s*:\s*(?P<value>{_flex(from_role)})(?=\s*(?:[.;,)]|$))")
    _fix_paragraph_lines(doc, pattern, to_role, fix)
    _fix_tables(doc, from_role, to_role, fix)

    if not fix.count:
        return fix

    # ── proof: only the role changed, and nothing was added or removed ──
    after = _texts(doc)
    if len(after) != len(before) or (len(doc.paragraphs), len(doc.tables)) != structure:
        raise RoleFixError("The change would have altered the document's structure, so it was not made.")
    for was, now in zip(before, after):
        if was != now and not _derives(was, now, from_role, to_role):
            raise RoleFixError("The change would have altered more than the role's name, so it was not made.")
    out = io.BytesIO()
    doc.save(out)
    fix.data = out.getvalue()
    return fix


def _derives(was: str, now: str, from_role: str, to_role: str) -> bool:
    """`now` is `was` with SOME of the occurrences of the old role replaced by the new one — and
    nothing else different. (A paragraph may name the role in a label and again in running text;
    only the label's is changed.)"""
    pieces = re.split(_flex(from_role), was)
    if len(pieces) < 2 or len(pieces) > 40:
        return False
    either = f"(?:{_flex(from_role)}|{re.escape(to_role)})"
    return re.fullmatch(either.join(re.escape(piece) for piece in pieces), now) is not None
